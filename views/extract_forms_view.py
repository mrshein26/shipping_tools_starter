import streamlit as st
import fitz  # PyMuPDF
import io
import os
import re
import zipfile
import pandas as pd
import textwrap
import subprocess
import tempfile
import shutil
from datetime import datetime, timedelta
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from pypdf import PdfReader, PdfWriter

# 💡 Excel Forms များအတွက် လိုအပ်သော Imports များ
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.drawing.image import Image as ExcelImage

from core.utils import add_log, extract_files

# =========================================================
# HELPER: PDF COMPRESSION
# =========================================================
def compress_pdf_bytes(pdf_bytes, gs_setting="/ebook"):
    """Ghostscript အသုံးပြု၍ PDF Bytes များကို Compress လုပ်ပေးသော Helper Function"""
    if not shutil.which("gs"):
        return pdf_bytes 
        
    with tempfile.TemporaryDirectory() as temp_dir:
        input_path = os.path.join(temp_dir, "input.pdf")
        output_path = os.path.join(temp_dir, "output.pdf")
        
        with open(input_path, "wb") as f:
            f.write(pdf_bytes)
            
        gs_cmd = [
            "gs", "-sDEVICE=pdfwrite", "-dCompatibilityLevel=1.4",
            f"-dPDFSETTINGS={gs_setting}", "-dNOPAUSE", "-dQUIET", "-dBATCH",
            f"-sOutputFile={output_path}", input_path
        ]
        
        try:
            subprocess.run(gs_cmd, check=True)
            if os.path.exists(output_path):
                with open(output_path, "rb") as f:
                    return f.read()
        except Exception:
            pass
            
    return pdf_bytes


# =========================================================
# EXTRACT FORMS LOGIC FUNCTIONS
# =========================================================

# --- 1. NWPM (CN) ---
def extract_invoice_data_for_dec(pdf_bytes):
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    text = "".join([page.get_text("text") for page in doc])
    clean_text = text.replace('\n', ' ')

    inv_m = re.search(r"Number\s*:\s*([A-Z0-9\-]+)", clean_text, re.IGNORECASE) or re.search(r"(MYHM[\d\-]+)", clean_text)
    inv_no = inv_m.group(1).strip() if inv_m else f"INV_{datetime.now().strftime('%H%M%S')}"
        
    date_m = re.search(r"Date\s*:\s*(\d{4}-\d{2}-\d{2})", clean_text, re.IGNORECASE)
    date_val = ""
    if date_m:
        try: 
            date_val = datetime.strptime(date_m.group(1), "%Y-%m-%d").strftime("%d.%m.%Y")
        except: 
            date_val = date_m.group(1)
        
    ord_m = re.search(r"Order No\s*:\s*([\d-]+)", clean_text, re.IGNORECASE) or re.search(r"(\d{6}-\d{4})", clean_text)
    ord_no = ord_m.group(1) if ord_m else ""
        
    ctn_m = re.search(r"(\d+)\s*Cartons?", clean_text, re.IGNORECASE)
    cartons = ctn_m.group(1) if ctn_m else ""
    
    qty_fallback = re.search(r"(\d+)\s+(?:\d+\.\d+)\s+(?:\d+\.\d+)", clean_text) or re.search(r"(\d+)\s*Pieces", clean_text, re.IGNORECASE)
    qty = qty_fallback.group(1) if qty_fallback else ""
        
    desc_block_match = re.search(r"Description of Goods(.*?)HS Code", clean_text, re.IGNORECASE)
    description = ""
    comp_text = ""
    
    if desc_block_match:
        desc_block = desc_block_match.group(1)
        desc_m = re.search(r"Cartons?\s+(.*?)(?=\d+%)", desc_block, re.IGNORECASE) or re.search(r"Cartons?\s+(.*)", desc_block, re.IGNORECASE)
        if desc_m: 
            description = desc_m.group(1).strip()
            
        comp_matches = re.findall(r"\d+%\s*[A-Z]+(?:\s*[A-Z\(\)]+)*", desc_block)
        comp_text = " ".join([m.strip() for m in comp_matches])

    if description: 
        description = re.sub(r"(?i)\s*(Pieces|USD|CTNS).*$", "", description).strip()
        
    return {
        "DESCRIPTION": description, 
        "INVOICE_NO": inv_no, 
        "DATE": date_val, 
        "ORDER_NO": ord_no, 
        "QUANTITY": qty, 
        "CARTONS": cartons, 
        "COMPOSITION": comp_text
    }

def generate_declaration_pdf(template_path, extracted_data):
    doc = fitz.open(template_path)
    page = doc[0]
    page.insert_text((160, 160), "EXPORTER DECLARATION OF USE", fontsize=14, fontname="hebo")
    page.insert_text((290, 180), "OF", fontsize=14, fontname="hebo")
    page.insert_text((150, 200), "NON-WOOD PACKAGING MATERIALS", fontsize=14, fontname="hebo")
    
    para1 = "THIS IS TO CERTIFY THAT THE FOLLOWING SHIPMENT OF WEARING APPAREL DOES NOT USE ANY SOLID WOOD PACKING MATERIALS (SWPM), INCLUDING, BUT NOT LIMITED TO, PALLETS, BRACING, BLOCKING, CRATING, DUNNAGE, PACKING BLOCKS, DRUMS, CASES OR SKIDS."
    para2 = "THE CLOTHING IS PACKED IN CORRUGATED CARTONS TO PREVENT DAMAGE WHILE BEING SHIPPED TO H&M HENNES & MAURITZ INC."
    
    page.insert_textbox(fitz.Rect(70, 230, 525, 295), para1, fontsize=10, fontname="helv", align=3)
    page.insert_textbox(fitz.Rect(70, 305, 525, 345), para2, fontsize=10, fontname="helv", align=3)
    
    y_start = 365
    label_x = 70
    value_x = 240
    
    fields = [
        ("DESCRIPTION", extracted_data.get("DESCRIPTION", "")), 
        ("INVOICE NO", extracted_data.get("INVOICE_NO", "")),
        ("DATE", extracted_data.get("DATE", "")), 
        ("ORDER NO", extracted_data.get("ORDER_NO", "")),
        ("QUANTITY (PCS)", extracted_data.get("QUANTITY", "")), 
        ("NO OF CARTON (CTNS)", extracted_data.get("CARTONS", "")),
        ("COMPOSITION", extracted_data.get("COMPOSITION", ""))
    ]
    
    current_y = y_start
    for label, value in fields:
        page.insert_text((label_x, current_y), label, fontsize=11, fontname="helv")
        rect = fitz.Rect(value_x, current_y - 11, 525, current_y + 30)
        page.insert_textbox(rect, f":   {value}", fontsize=11, fontname="helv", align=0)
        current_y += 38 if len(str(value)) > 55 else 25
        
    output_pdf = io.BytesIO()
    doc.save(output_pdf)
    doc.close()
    return output_pdf.getvalue()

# --- 2. DMAM (JP) ---
def extract_invoice_data_for_dmam(pdf_bytes):
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    full_text = "".join([page.get_text() for page in doc])
    order_no = ""
    order_match = re.search(r"H&M Order No:\s*([\w-]+)", full_text) or re.search(r"HM Order No:\s*([\w-]+)", full_text)
    if order_match: 
        order_no = order_match.group(1).strip()

    composition = ""
    block_match = re.search(r"Description of Goods(.*?)(?:HS Code|Container No|Total)", full_text, re.IGNORECASE | re.DOTALL)
    
    if block_match:
        perc_start = re.search(r"(\d+%[\s\S]+?)(?=\n\s*Pieces|\n\s*Pcs|\n\s*Sets|\n\s*\d+\s*\n|\Z)", block_match.group(1), re.IGNORECASE)
        if perc_start:
            lines = [l.strip() for l in perc_start.group(1).split('\n') if l.strip()]
            comp_lines = [l for l in lines if '%' in l or l.isupper() or l.isalpha()]
            composition = re.sub(r'\s+', ' ', " ".join(comp_lines).upper()).strip()
            
    if composition:
        for split_word in ['"', "THE SCIENTIFIC", "DOMESTICA"]:
            if split_word in composition: 
                composition = composition.split(split_word)[0].strip()
        composition = composition.rstrip(' ".,:')
        
    return order_no, composition

def generate_dmam_pdf(template_path, order_no, composition):
    doc = fitz.open(template_path)
    page = doc[0]
    font_reg = "helv"
    font_bold = "hebo"
    
    page.insert_text(fitz.Point(145, 125), "Declaration of Main and auxiliary material", fontsize=15, fontname=font_bold)
    page.draw_line(fitz.Point(145, 129), fitz.Point(446, 129), width=0.8)
    
    page.insert_text(fitz.Point(50, 165), "It is declared that this shipment", fontsize=11, fontname=font_reg)
    page.insert_text(fitz.Point(50, 185), f"Order No: {order_no}", fontsize=11, fontname=font_reg)
    page.insert_text(fitz.Point(50, 200), "Containing materials as below :", fontsize=11, fontname=font_reg)
    page.insert_text(fitz.Point(50, 235), "1, Main fabric information", fontsize=11, fontname=font_reg)
    page.insert_text(fitz.Point(50, 250), f"Composition = {composition}", fontsize=11, fontname=font_reg)
    page.insert_text(fitz.Point(50, 285), "the origin country at each point in each production process of all the fabric.", fontsize=11, fontname=font_reg)
    page.insert_text(fitz.Point(50, 300), 'B. Please add column if there are several "Other Fabric" (e.g. pocket and Cuffs,etc).', fontsize=11, fontname=font_reg)
    
    page.draw_rect(fitz.Rect(50, 315, 500, 435), width=1) 
    page.draw_line(fitz.Point(320, 315), fitz.Point(320, 435), width=1)
    
    for y_line in [335, 375, 395, 415]: 
        page.draw_line(fitz.Point(50, y_line), fitz.Point(500, y_line), width=1)
    
    page.insert_text(fitz.Point(325, 330), "Main", fontsize=11, fontname=font_reg)
    page.insert_text(fitz.Point(55, 350), "Origin Country of Raw material of Yarn", fontsize=11, fontname=font_reg)
    page.insert_text(fitz.Point(55, 365), "(e.g. Cotton, Polyester)", fontsize=11, fontname=font_reg)
    page.insert_text(fitz.Point(325, 360), "China", fontsize=11, fontname=font_reg)
    page.insert_text(fitz.Point(55, 390), "Country of Yarn Production", fontsize=11, fontname=font_reg)
    page.insert_text(fitz.Point(325, 390), "China", fontsize=11, fontname=font_reg)
    page.insert_text(fitz.Point(55, 410), "Country of Fabric Production", fontsize=11, fontname=font_reg)
    page.insert_text(fitz.Point(325, 410), "Myanmar", fontsize=11, fontname=font_reg)
    page.insert_text(fitz.Point(55, 430), "Country of where the fabric was Cut", fontsize=11, fontname=font_reg)
    page.insert_text(fitz.Point(325, 430), "Myanmar", fontsize=11, fontname=font_reg)
    
    page.insert_text(fitz.Point(50, 470), "3, Accessories information", fontsize=11, fontname=font_reg)
    page.insert_text(fitz.Point(50, 490), "Please specify all accessories/other which was imported from anywhere outside the origin which", fontsize=11, fontname=font_reg)
    page.insert_text(fitz.Point(50, 505), "the final product was made.", fontsize=11, fontname=font_reg)
    page.insert_text(fitz.Point(50, 525), "(e.g., ・ Sewing thread = MYANMAR", fontsize=11, fontname=font_reg)
    page.insert_text(fitz.Point(85, 540), "・ Label = MYANMAR", fontsize=11, fontname=font_reg)
    page.insert_text(fitz.Point(85, 555), "And so on)", fontsize=11, fontname=font_reg)

    output_pdf = io.BytesIO()
    doc.save(output_pdf)
    doc.close()
    return output_pdf.getvalue()

# --- 3. Org Criterion (IN) ---
def generate_org_criterion_pdf(template_bytes, new_order_no):
    doc = fitz.open(stream=template_bytes, filetype="pdf")
    page = doc[0]
    target_text = "562538-5878"
    text_instances = page.search_for(target_text)
    
    if text_instances:
        for inst in text_instances:
            page.draw_rect(fitz.Rect(inst.x0, inst.y0, inst.x1, inst.y1), color=(1, 1, 1), fill=(1, 1, 1))
            page.insert_text(fitz.Point(inst.x0, inst.y1 - 1), new_order_no, fontsize=6, fontname="helv", color=(0, 0, 0))
    else:
        page.insert_text(fitz.Point(72, 105), new_order_no, fontsize=6, fontname="helv", color=(0, 0, 0))
        
    out_pdf = io.BytesIO()
    doc.save(out_pdf)
    doc.close()
    return out_pdf.getvalue()

# --- 4. CO (CL) ---
def extract_data_for_cl(pdf_bytes):
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    text = "".join([page.get_text("text") for page in doc])
    clean_text = text.replace('\n', ' ')

    inv_m = re.search(r"Number\s*:\s*([A-Z0-9\-]+)", clean_text, re.IGNORECASE)
    inv_no = inv_m.group(1).strip() if inv_m else f"INV_{datetime.now().strftime('%H%M%S')}"
    
    desc_block_match = re.search(r"Description of Goods(.*?)HS Code", clean_text, re.IGNORECASE)
    description = "KNITTED GARMENTS"
    composition = ""
    
    if desc_block_match:
        desc_block = desc_block_match.group(1)
        comp_matches = re.findall(r"\d+%\s*[A-Z]+(?:\s*[A-Z\(\)]+)*", desc_block, re.IGNORECASE)
        composition = " ".join([m.strip().upper() for m in comp_matches])
        desc_m = re.search(r"Cartons?\s*(.*?)(?=\d+%|Pieces|Packs)", desc_block, re.IGNORECASE)
        if desc_m: 
            description = desc_m.group(1).strip().upper()
    
    hs_m = re.search(r"HS Code:\s*(\d{6})", clean_text, re.IGNORECASE)
    gw_m = re.search(r"Gross Weight:\s*([\d\.]+)", clean_text, re.IGNORECASE)
    date_m = re.search(r"Date\s*:\s*(\d{4}-\d{2}-\d{2})", clean_text, re.IGNORECASE)
    
    return {
        "CERT_NO": inv_no, 
        "DESCRIPTION": f"{description}\n{composition}".strip(), 
        "HS_CODE": hs_m.group(1) if hs_m else "", 
        "GROSS_WEIGHT": gw_m.group(1) if gw_m else "", 
        "DATE": date_m.group(1) if date_m else ""
    }

def generate_cl_pdf(template_path, data):
    doc = fitz.open(template_path)
    page = doc[0] 
    page.insert_text(fitz.Point(192, 209), data.get("CERT_NO", ""), fontsize=11, fontname="helv")
    page.insert_textbox(fitz.Rect(70, 400, 260, 520), data.get("DESCRIPTION", ""), fontsize=10, fontname="helv", align=0)
    page.insert_text(fitz.Point(315, 410), data.get("HS_CODE", ""), fontsize=10, fontname="helv")
    page.insert_text(fitz.Point(490, 410), data.get("GROSS_WEIGHT", ""), fontsize=10, fontname="helv")
    
    place_and_date = f"Yangon, {data.get('DATE', '')}"
    page.insert_text(fitz.Point(120, 682), place_and_date, fontsize=10, fontname="helv")
    page.insert_text(fitz.Point(392, 708), place_and_date, fontsize=10, fontname="helv")
    
    out_pdf = io.BytesIO()
    doc.save(out_pdf)
    doc.close()
    return out_pdf.getvalue()

# --- 5. Cover Sheet ---
def generate_cover_sheets(excel_file, template_path, do_compress=False):
    zip_buffer = io.BytesIO()
    count = 0
    try: 
        df = pd.read_excel(excel_file)
    except: 
        excel_file.seek(0)
        df = pd.read_csv(excel_file)

    with open(template_path, "rb") as f: 
        template_bytes = f.read()

    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for index, row in df.iterrows():
            po_no = str(row.get('PO NO', str(row.iloc[0]))).strip().split('.')[0]
            sku = str(row.get('SKU', '')).strip()
            wh_code = str(row.get('WH CODE', '')).strip()
            if not po_no or po_no.lower() == 'nan': 
                continue

            try:
                doc = fitz.open(stream=template_bytes, filetype="pdf")
                page = doc[0]
                if sku and sku.lower() != 'nan':
                    page.insert_text(fitz.Point(140, 123), sku, fontsize=11, fontname="hebo")
                page.insert_text(fitz.Point(140, 144), po_no, fontsize=11, fontname="hebo")
                
                pdf_bytes = io.BytesIO()
                doc.save(pdf_bytes)
                doc.close()
                
                final_bytes = pdf_bytes.getvalue()
                if do_compress:
                    final_bytes = compress_pdf_bytes(final_bytes)
                    
                zf.writestr(f"{po_no} {sku}-{wh_code}.pdf", final_bytes)
                count += 1
            except Exception as e:
                pass
                
    return zip_buffer.getvalue() if count > 0 else None, count


# =========================================================
# MAIN VIEW RENDER
# =========================================================
def render_extract_forms_ui():
    st.subheader("📝 Extract Forms")
    
    if "form_up_key" not in st.session_state: 
        st.session_state.form_up_key = 0
    if "air_bkg_res" not in st.session_state:
        st.session_state.air_bkg_res = None
        
    def clear_files(): 
        st.session_state.form_up_key += 1
        st.session_state.air_bkg_res = None

    TEMPLATE_DIR = "assets/templates"
    if not os.path.exists(TEMPLATE_DIR): 
        os.makedirs(TEMPLATE_DIR)

    # 🗂️ ဖောင်အမျိုးအစား ရွေးချယ်ခြင်း (Air Booking Excel ကို ထပ်မံထည့်သွင်းထားသည်)
    st.write("Please select the required Form Type:")
    selected_form = st.radio(
        "Select Form Type",
        ["Declaration of NWPM (CN)", "Statement of Origin (CA,DR)", "DMAM (JP)", "Org Criterion (IN)", "Cover Sheet", "Certificate of Origin (CL)", "Air Booking (Excel)"],
        horizontal=True, label_visibility="collapsed", key="form_selector"
    )
    
    st.markdown("<br>", unsafe_allow_html=True) # ခပ်ပါးပါး ခြားရန်
    
    do_compress = st.toggle("🗜️ **Compress Output PDFs**", value=True)
    if do_compress and not shutil.which("gs"):
        st.warning("⚠️ သင့်စက်တွင် Ghostscript ကို Install မလုပ်ထားပါ။ Compression အလုပ်လုပ်မည် မဟုတ်ပါ။")
        
    st.markdown("---") # အောက်ပိုင်းနှင့်ပိုင်းခြားရန် တစ်ကြောင်းသာ ထားမည်

    # [1] Declaration of NWPM (CN)
    if selected_form == "Declaration of NWPM (CN)":
        st.markdown("#### 📄 Declaration of Non Wood Packaging Materials")
        st.info("Upload CN, OB Invoice PDFs to extract the Forms.")
        inv_files = st.file_uploader("Upload Invoice PDFs", type=["pdf", "zip"], accept_multiple_files=True, key=f"cn_{st.session_state.form_up_key}")
        
        c1, c2 = st.columns(2)
        if c1.button("⚡ Generate Forms", use_container_width=True, type="primary", key="btn_ex_dec"):
            if not inv_files: 
                st.error("❌ Please upload Invoice PDF files first.")
            else:
                with st.spinner("Generating Exporter Declarations..."):
                    template_file = os.path.join(TEMPLATE_DIR, "Declaration_Template.pdf")
                    if not os.path.exists(template_file): 
                        st.error(f"❌ Master template file `{template_file}` not found!")
                    else:
                        extracted_files = extract_files(inv_files)
                        zip_buffer = io.BytesIO()
                        success_count = 0
                        
                        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
                            for f_obj in extracted_files:
                                try:
                                    f_obj.seek(0)
                                    data = extract_invoice_data_for_dec(f_obj.read())
                                    generated_pdf = generate_declaration_pdf(template_file, data)
                                    if generated_pdf:
                                        if do_compress:
                                            generated_pdf = compress_pdf_bytes(generated_pdf)
                                        zf.writestr(f"{os.path.splitext(f_obj.name)[0]}_NWPM.pdf", generated_pdf)
                                        success_count += 1
                                except Exception as e: 
                                    st.error(f"Error processing {f_obj.name}: {e}")
                                    
                        if success_count > 0:
                            st.success(f"✅ Successfully generated {success_count} Declaration Forms!")
                            add_log(st.session_state.current_user, f"Extract Forms: Generated {success_count} NWPM(CN)")
                            st.download_button("📥 Download Generated Forms (ZIP)", zip_buffer.getvalue(), "Exporter_Declarations_NWPM(CN).zip", "application/zip", use_container_width=True)
                            
        if c2.button("🗑️ Clear Files", use_container_width=True, key="clr_ex_dec"): 
            clear_files()
            st.rerun()

    # [2] Statement of Origin (CA,DR)
    elif selected_form == "Statement of Origin (CA,DR)":
        st.markdown("#### 📝 Statement of Origin, Transport Document, Certificate of Origin, Canada Customs Invoice")
        st.info("Upload CA,DR Invoice PDFs to extract the Forms.")
        
        template_file = os.path.join(TEMPLATE_DIR, "CA Template.pdf")
        if not os.path.exists(template_file):
            st.error(f"❌ Error: '{template_file}' template ကို `assets/templates/` ထဲတွင် ရှာမတွေ့ပါ။")
        else:
            with open(template_file, "rb") as f: 
                ca_template_bytes = f.read()
                
            ca_inv_files = st.file_uploader("Upload Invoice PDFs", type=["pdf", "zip"], accept_multiple_files=True, key=f"ca_inv_{st.session_state.form_up_key}")
            
            c_ca1, c_ca2 = st.columns(2)
            if c_ca1.button("⚡ Generate Forms", key="btn_gen_ca", use_container_width=True, type="primary"):
                if not ca_inv_files: 
                    st.error("⚠️ ကျေးဇူးပြု၍ Invoice PDF ဖိုင်များကို အရင် Upload တင်ပေးပါ။")
                else:
                    with st.spinner("Processing All Pages of Statement & Certificate Forms... (Compression may take a while)"):
                        try:
                            extracted_ca_files = extract_files(ca_inv_files)
                            zip_buffer = io.BytesIO()
                            success_count = 0
                            
                            with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
                                for f_obj in extracted_ca_files:
                                    try:
                                        f_obj.seek(0)
                                        doc = fitz.open(stream=f_obj.read(), filetype="pdf")
                                        full_text = "".join(page.get_text("text") for page in doc)
                                        clean_text = full_text.replace('\n', ' ')
                                        
                                        inv_num_m = re.search(r"Number\s*:\s*([A-Z0-9\-]+)", clean_text, re.IGNORECASE) or re.search(r"Number:\s*(MYHM\d+)", clean_text, re.IGNORECASE)
                                        inv_date_m = re.search(r"Date\s*:\s*(\d{4}-\d{2}-\d{2})", clean_text, re.IGNORECASE) or re.search(r"Date:\s*(\d{4}-\d{2}-\d{2})", clean_text, re.IGNORECASE)
                                        order_m = re.search(r"H&M Order No\s*:\s*([\d\-]+)", clean_text, re.IGNORECASE) or re.search(r"(\d{6}-\d{4})", clean_text)
                                        
                                        inv_no_val = inv_num_m.group(1).strip() if inv_num_m else "N/A"
                                        inv_date_val = inv_date_m.group(1).strip() if inv_date_m else "N/A"
                                        inv_no_date = f"{inv_no_val} ({inv_date_val})"
                                        order_no = order_m.group(1).strip() if order_m else "N/A"
                                        
                                        desc_val = "KNITTED WEAR"
                                        comp_val = "100% COTTON"
                                        block_match = re.search(r"Description of Goods(.*?)(?:HS Code|Container No|Total)", clean_text, re.IGNORECASE | re.DOTALL)
                                        
                                        if block_match:
                                            desc_block = block_match.group(1)
                                            desc_match = re.search(r"(Cardigan|Polo shirt|Jumper|Shirt|Sweater|T-shirt|Trousers|Shorts)\s+[A-Za-z\s]+Knitted", desc_block, re.IGNORECASE)
                                            if desc_match: 
                                                desc_val = desc_match.group(0).strip()
                                                
                                            clean_desc_block = re.sub(r"\s+", " ", desc_block).strip()
                                            raw_matches = re.findall(r"(\d+\s*%\s*[A-Z\(\)\-]+)", clean_desc_block, re.IGNORECASE)
                                            comp_parts = []
                                            for match in raw_matches:
                                                clean_part = re.sub(r"(Cardigan|Polo shirt|Jumper|Shirt|Sweater|T-shirt|Trousers|Shorts|Knitted)", "", match.strip(), flags=re.IGNORECASE).strip()
                                                if clean_part and clean_part != "%": 
                                                    comp_parts.append(clean_part)
                                            if comp_parts: 
                                                comp_val = re.sub(r"\s+", " ", " ".join(comp_parts)).strip()

                                        gw_match = re.search(r"Gross Weight:\s*([\d\.\s]+(?:KG|kg)?)", clean_text, re.IGNORECASE)
                                        gw_val = gw_match.group(1).strip().upper() if gw_match else "N/A"
                                        pkgs_match = re.search(r"(\d+)\s*Cartons", clean_text, re.IGNORECASE)
                                        pkgs_count = pkgs_match.group(1).strip() if pkgs_match else "N/A"
                                        pkgs_display = f"{pkgs_count} CTNS" if pkgs_count != "N/A" else "N/A"
                                        
                                        try: 
                                            formatted_date = datetime.strptime(inv_date_val, "%Y-%m-%d").strftime("%d/%m/%Y")
                                        except: 
                                            formatted_date = inv_date_val

                                        nw_match_cci = re.search(r"Net Weight:\s*([\d\.\s]+(?:KG|kg)?)", clean_text, re.IGNORECASE)
                                        nw_val_cci = nw_match_cci.group(1).strip().upper() if nw_match_cci else ""
                                        hs_match_cci = re.search(r"HS Code:\s*([\d]+)", clean_text, re.IGNORECASE)
                                        hs_val_cci = hs_match_cci.group(1).strip() if hs_match_cci else ""

                                        qty_val_cci, unit_price_cci, total_amt_cci = "", "", ""
                                        raw_ints, raw_decs = [], []
                                        for word in clean_text.split():
                                            word = word.strip('()[]:,;$')
                                            if re.match(r"^[1-9][0-9,]*$", word): 
                                                raw_ints.append(word.replace(',', ''))
                                            elif re.match(r"^[0-9,]+\.[0-9]{2}$", word): 
                                                raw_decs.append(word.replace(',', ''))
                                                
                                        all_ints = sorted(list(set([int(x) for x in raw_ints])), reverse=True)
                                        all_decs = list(set([float(x) for x in raw_decs]))
                                        
                                        tot_m = re.search(r"(?:Total after Discount|Total|Amount Chargeable)[\sA-Za-z\:]*?([0-9,]+\.[0-9]{2})", clean_text, re.IGNORECASE)
                                        explicit_t_val = float(tot_m.group(1).replace(',', '')) if tot_m else None

                                        found_math = False
                                        if explicit_t_val and explicit_t_val in all_decs:
                                            for q_val in all_ints:
                                                for p_val in all_decs:
                                                    if q_val > 0 and p_val > 0 and abs((q_val * p_val) - explicit_t_val) <= 0.1:
                                                        qty_val_cci, unit_price_cci, total_amt_cci = str(q_val), f"{p_val:.2f}", f"{explicit_t_val:.2f}"
                                                        found_math = True
                                                        break
                                                if found_math: break
                                                
                                        if not found_math:
                                            for q_val in all_ints:
                                                if q_val < 2: continue
                                                for p_val in all_decs:
                                                    for t_val in all_decs:
                                                        if t_val > p_val and p_val > 0 and abs((q_val * p_val) - t_val) <= 0.1:
                                                            qty_val_cci, unit_price_cci, total_amt_cci = str(q_val), f"{p_val:.2f}", f"{t_val:.2f}"
                                                            found_math = True
                                                            break
                                                    if found_math: break
                                                if found_math: break
                                                
                                        if not found_math:
                                            qty_match_cci = re.search(r"(?:Quantity|Qty|Pieces|Pcs)[^a-zA-Z0-9]*([0-9,]+)", clean_text, re.IGNORECASE) or re.search(r"([0-9,]+)[^a-zA-Z0-9]*(?:Pieces|Pcs)", clean_text, re.IGNORECASE)
                                            qty_val_cci = qty_match_cci.group(1).strip() if qty_match_cci else ""
                                            prices_cci = re.findall(r"(?:USD|EUR|GBP|\$|CAD)?\s*([0-9,]+\.[0-9]{2})", clean_text, re.IGNORECASE)
                                            if len(prices_cci) >= 2: 
                                                unit_price_cci, total_amt_cci = prices_cci[0], prices_cci[-1]
                                            elif len(prices_cci) == 1: 
                                                unit_price_cci, total_amt_cci = prices_cci[0], prices_cci[0]

                                        qty_display = qty_val_cci
                                        try: 
                                            cci_date = datetime.strptime(inv_date_val, "%Y-%m-%d").strftime("%m/%d/%y")
                                        except: 
                                            cci_date = inv_date_val

                                        # Drawing
                                        packet = io.BytesIO()
                                        can = canvas.Canvas(packet, pagesize=letter)
                                        
                                        can.setFont("Helvetica-Bold", 12)
                                        can.drawCentredString(306, 610, "EXPORTER'S STATEMENT OF ORIGIN")
                                        can.setFont("Helvetica", 10)
                                        can.drawString(60, 580, "I CERTIFY THAT THE GOODS DESCRIBED IN THIS INVOICE OR IN THE ATTACHED INVOICE NO.")
                                        can.drawString(60, 565, "REPRODUCED IN THE BENEFICIARY COUNTRY OF ORIGIN AND THAT AT LEAST")
                                        can.drawString(80, 520, "INVOICE NO & DATE")
                                        can.drawString(205, 520, f": {inv_no_date.upper()}")
                                        can.drawString(80, 495, "DESCRIPTION")
                                        can.drawString(205, 495, f": {desc_val.upper()}")
                                        can.drawString(80, 470, "COMPOSITION")
                                        can.drawString(205, 470, f": {comp_val.upper()}")
                                        can.drawString(80, 445, "ORDER NO")
                                        can.drawString(205, 445, f": {order_no.upper()}")
                                        can.drawString(80, 410, "NAME AND TITLE")
                                        can.drawString(205, 410, ": JIANGSU CENTURY LIAOYUAN KNITTED WEAR CO., LTD")
                                        can.drawString(80, 385, "CORPORATION")
                                        can.drawString(205, 385, ": TENG HUI (MYANMAR) KNITTING COMPANY LIMITED")
                                        can.drawString(80, 360, "TELEPHONE NO")
                                        can.drawString(205, 360, ": 09420035004")
                                        
                                        # P2
                                        can.showPage()
                                        can.setFont("Helvetica", 9)
                                        can.drawString(400, 735, f"{inv_no_val.upper()}")
                                        can.drawString(400, 710, f"{formatted_date}")
                                        can.drawCentredString(40, 490, f"{pkgs_count}")
                                        can.drawString(100, 490, f"{desc_val.upper()}")
                                        y_pos_p2 = 478
                                        for line in textwrap.wrap(comp_val.upper(), width=60): 
                                            can.drawString(100, y_pos_p2, line)
                                            y_pos_p2 -= 12
                                        can.drawCentredString(515, 490, f"{gw_val}")
                                        
                                        # P3
                                        can.showPage()
                                        can.setFont("Helvetica", 8)
                                        can.drawCentredString(92, 510, f"{pkgs_display.upper()}")
                                        can.drawString(195, 510, f"{desc_val.upper()}")
                                        y_pos_p3 = 498
                                        for line in textwrap.wrap(comp_val.upper(), width=50): 
                                            can.drawString(195, y_pos_p3, line)
                                            y_pos_p3 -= 12
                                        can.drawCentredString(490, 510, f"{inv_no_val.upper()}")
                                        can.drawCentredString(490, 498, f"{formatted_date}")
                                        
                                        # P4
                                        can.showPage()
                                        can.setFont("Helvetica", 9)
                                        can.drawString(285, 770, f"{cci_date}")
                                        can.drawString(285, 740, f"{inv_no_val.upper()}")
                                        can.drawString(285, 728, f"{order_no}")
                                        can.drawCentredString(40, 490, f"{pkgs_count}")
                                        can.drawString(60, 490, f"{desc_val.upper()}")
                                        
                                        y_pos_p4 = 478
                                        max_len = 45
                                        comp_str = comp_val.upper()
                                        if len(comp_str) > max_len:
                                            split_idx = comp_str.rfind(' ', 0, max_len)
                                            if split_idx == -1: split_idx = max_len
                                            can.drawString(60, y_pos_p4, comp_str[:split_idx])
                                            y_pos_p4 -= 12
                                            can.drawString(60, y_pos_p4, comp_str[split_idx:].strip())
                                        else: 
                                            can.drawString(60, y_pos_p4, comp_str)
                                            
                                        can.drawString(60, y_pos_p4 - 17, f"HS Code: {hs_val_cci}")
                                        can.drawCentredString(427, 490, f"{qty_display or '000'}")
                                        can.drawCentredString(492, 490, f"{unit_price_cci or '00.00'}")
                                        if total_amt_cci: 
                                            can.drawCentredString(555, 490, f"{total_amt_cci}")
                                            
                                        can.drawCentredString(427, 325, f"{nw_val_cci.replace('KG', '').replace('kg', '').strip()}")
                                        can.drawCentredString(488, 325, f"{gw_val.replace('KG', '').replace('kg', '').strip()}")
                                        can.drawCentredString(555, 325, f"{total_amt_cci}")
                                        can.drawString(130, 325, f"{inv_no_val.upper()}")
                                        
                                        can.save()
                                        packet.seek(0)
                                        
                                        template_pdf = PdfReader(io.BytesIO(ca_template_bytes))
                                        new_pdf = PdfReader(packet)
                                        output_writer = PdfWriter()
                                        
                                        for idx in range(len(template_pdf.pages)):
                                            t_page = template_pdf.pages[idx]
                                            if idx < len(new_pdf.pages): 
                                                t_page.merge_page(new_pdf.pages[idx])
                                            output_writer.add_page(t_page)
                                            
                                        final_pdf_bytes = io.BytesIO()
                                        output_writer.write(final_pdf_bytes)
                                        final_pdf_bytes.seek(0)
                                        
                                        # 💡 Compression for CA, DR Forms
                                        final_data = final_pdf_bytes.read()
                                        if do_compress:
                                            final_data = compress_pdf_bytes(final_data)
                                        
                                        base_name = f_obj.name.replace(".pdf", "").replace(".PDF", "")
                                        zf.writestr(f"{base_name} SO-TR-CO-CCI.pdf", final_data)
                                        success_count += 1
                                        
                                    except Exception as e:
                                        st.error(f"Error processing {f_obj.name}: {e}")
                                
                            if success_count > 0:
                                st.success(f"✅ Successfully generated {success_count} CA Full Forms!")
                                add_log(st.session_state.current_user, f"Extract Forms: Generated {success_count} x CA Forms")
                                st.download_button("📥 Download CA Forms (ZIP)", zip_buffer.getvalue(), "CA_Forms.zip", "application/zip", use_container_width=True)
                                
                        except Exception as e: 
                            st.error(f"❌ Error: {e}")
                            
            if c_ca2.button("🗑️ Clear Files", key="clr_ca_forms", use_container_width=True): 
                clear_files()
                st.rerun()

    # [3] DMAM (JP)
    elif selected_form == "DMAM (JP)":
        st.markdown("#### 📄 Declaration of Main and Auxiliary Material")
        st.info("Upload JP, OJ Invoice PDFs to extract the forms.")
        dmam_files = st.file_uploader("Upload Invoice PDFs", type=["pdf", "zip"], accept_multiple_files=True, key=f"jp_{st.session_state.form_up_key}")
        
        c1, c2 = st.columns(2)
        if c1.button("⚡ Generate Forms", use_container_width=True, type="primary", key="btn_generate_dmam"):
            if not dmam_files: 
                st.error("❌ Please upload Invoice PDF files first.")
            else:
                with st.spinner("Generating DMAM (JP) Forms..."):
                    template_file = os.path.join(TEMPLATE_DIR, "DMAM_Template.pdf")
                    if not os.path.exists(template_file): 
                        st.error(f"❌ Master template file `{template_file}` not found!")
                    else:
                        extracted_files = extract_files(dmam_files)
                        zip_buffer = io.BytesIO()
                        success_count = 0
                        try:
                            with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
                                for f_obj in extracted_files:
                                    try:
                                        f_obj.seek(0)
                                        order_no, composition = extract_invoice_data_for_dmam(f_obj.read())
                                        gen_pdf = generate_dmam_pdf(template_file, order_no, composition)
                                        if gen_pdf:
                                            if do_compress:
                                                gen_pdf = compress_pdf_bytes(gen_pdf)
                                            base_name = os.path.splitext(f_obj.name)[0]
                                            if base_name.upper().endswith("_PL"): 
                                                base_name = base_name[:-3]
                                            zf.writestr(f"{base_name}_DMAM.pdf", gen_pdf)
                                            success_count += 1
                                    except: 
                                        pass
                                        
                            if success_count > 0:
                                st.success(f"🎉 Successfully generated {success_count} Dec M&A Forms!")
                                add_log(st.session_state.current_user, f"Extract Forms: Generated {success_count} x DMAM (JP) Forms")
                                st.download_button("📥 Download DMAM Forms (ZIP)", zip_buffer.getvalue(), "DMAM(JP)_Forms.zip", "application/zip", use_container_width=True)
                                
                        except Exception as e: 
                            st.error(f"Error processing files: {e}")
                            
        if c2.button("🗑️ Clear Files", use_container_width=True, key="clr_dmam_files"): 
            clear_files()
            st.rerun()

    # [4] Org Criterion (IN)
    elif selected_form == "Org Criterion (IN)":
        st.markdown("#### 📄 Org Criterion (IN OI)")
        st.info("Upload 'OC INOI Lists.xlsx' file to generate the Forms.")
        oc_excel_file = st.file_uploader("Upload Excel or CSV File", type=["xlsx", "xls", "csv"], key=f"in_{st.session_state.form_up_key}")
        
        c1, c2 = st.columns(2)
        generate_clicked = c1.button("⚡ Generate Forms", use_container_width=True, type="primary", key="btn_generate_oc")
        
        if c2.button("🗑️ Clear Files", use_container_width=True, key="clr_oc_files"): 
            clear_files()
            st.rerun()

        if generate_clicked:
            if not oc_excel_file: 
                st.error("❌ Please upload the Excel/CSV file first.")
            else:
                with st.spinner("Generating Org Criterion Forms..."):
                    template_file = os.path.join(TEMPLATE_DIR, "Originating Criterion (IN OI).pdf")
                    if not os.path.exists(template_file): 
                        st.error(f"❌ Master template file `{template_file}` not found!")
                    else:
                        try:
                            with open(template_file, "rb") as f: 
                                template_bytes = f.read()
                                
                            try: 
                                df = pd.read_excel(oc_excel_file)
                            except: 
                                oc_excel_file.seek(0)
                                df = pd.read_csv(oc_excel_file)
                            
                            zip_buffer = io.BytesIO()
                            success_count = 0
                            
                            with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
                                for index, row in df.iterrows():
                                    po_no = str(row.get('PO NO', '')).strip().split('.')[0]
                                    depts = str(row.get('Depts', '')).strip().split('.')[0]
                                    sku = str(row.get('SKU', '')).strip()
                                    wh_code = str(row.get('WH CODE', '')).strip()
                                    
                                    if not po_no or po_no.lower() == 'nan': 
                                        continue
                                    
                                    order_no = f"{po_no}-{depts}" if depts and depts.lower() != 'nan' else po_no
                                    gen_pdf = generate_org_criterion_pdf(template_bytes, order_no)
                                    
                                    if gen_pdf:
                                        if do_compress:
                                            gen_pdf = compress_pdf_bytes(gen_pdf)
                                        zf.writestr(f"{po_no} {sku}-{wh_code}.pdf", gen_pdf)
                                        success_count += 1
                                        
                            if success_count > 0:
                                st.success(f"🎉 Successfully generated {success_count} Org Criterion Forms!")
                                add_log(st.session_state.current_user, f"Extract Forms: Generated {success_count} x Org Criterion Forms")
                                st.download_button("📥 Download Org Criterion Forms (ZIP)", zip_buffer.getvalue(), "Org_Criterion_Forms(IN).zip", "application/zip", use_container_width=True)
                                
                        except Exception as e: 
                            st.error(f"Error processing files: {e}")

    # [5] Cover Sheet
    elif selected_form == "Cover Sheet":
        st.markdown("### 📄 Cover Sheet")        
        st.info("Upload CV Sheet.xlsx to generate the Forms (PA,GB,OG,IX,OD,ME,IDW224)")
        uploaded_excel = st.file_uploader("Upload Excel List", type=["xlsx", "xls", "csv"], key=f"cover_{st.session_state.form_up_key}")
        
        c1, c2 = st.columns(2)
        generate_clicked = c1.button("⚡ Generate Cover Sheet", use_container_width=True, type="primary", key="btn_gen_cover")
        
        if c2.button("🗑️ Clear Files", use_container_width=True, key="clear_cover_sheet"): 
            clear_files()
            st.rerun()

        if generate_clicked:
            if not uploaded_excel: 
                st.error("❌ Please upload the Excel/CSV file first.")
            else:
                with st.spinner("Generating Cover Sheets..."):
                    template_path = os.path.join(TEMPLATE_DIR, "Cover_Sheet_Template.pdf")
                    if not os.path.exists(template_path): 
                        st.error(f"❌ Master template file `{template_path}` not found!")
                    else:
                        zip_file_buffer, total_count = generate_cover_sheets(uploaded_excel, template_path, do_compress)
                        if zip_file_buffer and total_count > 0:
                            st.success(f"✅ {total_count} Cover Sheets generated successfully!")
                            add_log(st.session_state.current_user, f"Extract Forms: Generated {total_count} x Cover Sheets")
                            st.download_button("📥 Download All Cover Sheets (ZIP)", zip_file_buffer, "Generated_Cover_Sheets.zip", "application/zip", use_container_width=True)

    # [6] Certificate of Origin (CL)
    elif selected_form == "Certificate of Origin (CL)":
        st.markdown("#### 📄 Certificado De Origen (Chile)")
        st.info("Upload CL Invoice PDFs to automatically generate Certificate of Origin Forms.")
        cl_files = st.file_uploader("Upload Invoice PDFs", type=["pdf", "zip"], accept_multiple_files=True, key=f"cl_{st.session_state.form_up_key}")
        
        c1, c2 = st.columns(2)
        if c1.button("⚡ Generate Forms", use_container_width=True, type="primary", key="btn_generate_cl"):
            if not cl_files: 
                st.error("❌ Please upload Invoice PDF files first.")
            else:
                with st.spinner("Generating Certificado De Origen (CL) Forms..."):
                    template_file = os.path.join(TEMPLATE_DIR, "CERTIFICADO DE ORIGEN (CL).pdf") 
                    if not os.path.exists(template_file): 
                        st.error(f"❌ Master template file `{template_file}` not found!")
                    else:
                        extracted_files = extract_files(cl_files)
                        zip_buffer = io.BytesIO()
                        success_count = 0
                        try:
                            with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
                                for f_obj in extracted_files:
                                    try:
                                        f_obj.seek(0)
                                        extracted_data = extract_data_for_cl(f_obj.read())
                                        gen_pdf = generate_cl_pdf(template_file, extracted_data)
                                        if gen_pdf:
                                            if do_compress:
                                                gen_pdf = compress_pdf_bytes(gen_pdf)
                                            zf.writestr(f"{os.path.splitext(f_obj.name)[0]}_CO.pdf", gen_pdf)
                                            success_count += 1
                                    except: 
                                        pass
                                        
                            if success_count > 0:
                                st.success(f"🎉 Successfully generated {success_count} Certificate of Origin (CL) Forms!")
                                add_log(st.session_state.current_user, f"Extract Forms: Generated {success_count} x Certificado De Origen (CL)")
                                st.download_button("📥 Download Generated CL Forms (ZIP)", zip_buffer.getvalue(), "Certificate_of_Origin_CL_Forms.zip", "application/zip", use_container_width=True)
                                
                        except Exception as e: 
                            st.error(f"Error processing files: {e}")
                            
        if c2.button("🗑️ Clear Files", use_container_width=True, key="clr_cl_files"): 
            clear_files()
            st.rerun()

    # 🛩️ [7] Air Booking (Excel) အသစ်ထည့်သွင်းခြင်း
    elif selected_form == "Air Booking (Excel)":
        st.markdown("#### ✈️ Air Booking Automation")
        st.info("Upload **H&M Air BKG Lists.xlsx** to generate automated Air Booking Excel forms.")
        bkg_file = st.file_uploader("Upload H&M Air BKG Lists (Excel)", type=["xlsx"], key=f"air_bkg_{st.session_state.form_up_key}")
        
        c_air1, c_air2 = st.columns(2)
        if c_air1.button("⚡ Generate Forms", use_container_width=True, type="primary", key="btn_air_gen"):
            if not bkg_file: 
                st.error("❌ ကျေးဇူးပြု၍ BKG List ဖိုင်ကို Upload တင်ပေးပါ။")
            else:
                with st.spinner("Generating Air Booking Excel files..."):
                    address_file = os.path.join(TEMPLATE_DIR, 'Addresses for HAWB [BKG] - 15.10.2025.xlsx')
                    logo_image_file = 'assets/images/logo.png'
                    stamp_image_file = 'assets/images/stamp.png' 
                    
                    if not os.path.exists(address_file):
                        st.error(f"❌ Error: '{address_file}' ကို Templates ဖိုဒါထဲတွင် ရှာမတွေ့ပါ။")
                    else:
                        try:
                            current_date = datetime.now().strftime("%d.%m.%y")
                            xl = pd.ExcelFile(bkg_file)
                            df_bkg = pd.read_excel(bkg_file, sheet_name=0, skiprows=1)
                            df_addr = pd.read_excel(address_file, sheet_name='HAWB')
                            tod_date = xl.sheet_names[0]
                            
                            light_grey_side = Side(style='thin', color='B0B0B0')
                            full_border = Border(left=light_grey_side, right=light_grey_side, top=light_grey_side, bottom=light_grey_side)
                            date_border = Border(right=light_grey_side, bottom=light_grey_side)

                            zip_buf = io.BytesIO()
                            file_count = 1
                            
                            with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
                                for index, row in df_bkg.iterrows():
                                    if pd.isna(row.get('WH CODE')): continue
                                    wh_code_bkg = str(row['WH CODE']).strip()
                                    order_no = str(row['PO NO']).split('.')[0] 
                                    depts_no = str(row['Depts']).split('.')[0] 
                                    country_code = str(row['SKU']).strip() if not pd.isna(row.get('SKU')) else ""
                                    
                                    mask = df_addr['Final Receiver'].str.contains(wh_code_bkg, na=False, case=False)
                                    addr_match = df_addr[mask]
                                    dest_name = addr_match.iloc[0]['Destination Country'] if not addr_match.empty else "NOT FOUND"
                                    dest_and_code = f"{dest_name} / {country_code}"
                                    consignee_val = addr_match.iloc[0]['HAWB Consignee '] if not addr_match.empty else "NOT FOUND"

                                    wb = Workbook()
                                    ws = wb.active
                                    ws.title = "Air Booking Form"
                                    ws.sheet_view.showGridLines = False

                                    ws.page_setup.paperSize = ws.PAPERSIZE_A4
                                    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
                                    ws.sheet_properties.pageSetUpPr.fitToPage = True
                                    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
                                    ws.page_margins.left = 0.25; ws.page_margins.right = 0.25
                                    ws.page_margins.top = 0.25; ws.page_margins.bottom = 0.25

                                    ws.column_dimensions['A'].width = 46; ws.column_dimensions['B'].width = 80
                                    ws.row_dimensions[1].height = 110
                                    ws.merge_cells('A1:B1')
                                    if os.path.exists(logo_image_file):
                                        ws.add_image(ExcelImage(logo_image_file), 'A1')
                                        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
                                        
                                    ws.row_dimensions[2].height = 25
                                    ws.merge_cells('A2:B2')
                                    ws['A2'] = f" DATE: {current_date}"
                                    ws['A2'].font, ws['A2'].alignment, ws['A2'].border = Font(size=10), Alignment(horizontal='left', vertical='center'), date_border

                                    content = [
                                        ("SHIPPER NAME AND ADDRESS", "TENG HUI (MYANMAR) KNITTING COMPANY LIMITED\nNO.8/KHA KWIN NO.300/A, RAKHINE YOE GYI VILLAGE TRACT,\nHTAN TA BIN TOWNSHIP, YANGON REGION,\nREPUBLIC OF THE UNION OF MYANMAR."),
                                        ("CONSIGNEE NAME AND ADDRESS", consignee_val), ("NOTIFY", "SAME AS CONSIGNEE"), ("TOD DATE", tod_date),
                                        ("H & M ORDER NUMBER", f"{order_no}/{depts_no}"), ("DESCRIPTION OF CARGO", f"{row.get('DESCRIPTION', '')}\n{row.get('COMPOSITION', '')}"), 
                                        ("NUMBER OF PIECES (PCS)", row.get('PCS', '')), ("NUMBER OF CARTONS (CTN)", row.get('CTNS', '')),
                                        ("TOTAL GROSS WEIGHT (KG)", row.get('G.W', '')), ("TOTAL VOLUME (CBM)", row.get('CBM', '')),
                                        ("MODE OF TRANSPORT", "BY AIR"), ("DESTINATION AND COUNTRY CODE", dest_and_code),
                                        ("FREIGHT TERM", "FREIGHT COLLECT"), ("SEASON AND DESTINATION WAREHOUSE", f"3-2026/{wh_code_bkg}"),
                                        ("H.S CODE", row.get('HS CODE', ''))
                                    ]

                                    current_row = 3
                                    for label, value in content:
                                        line_count = str(value).count('\n') + 1
                                        if "SHIPPER" in label: row_h = 85
                                        elif "CONSIGNEE" in label or "DESCRIPTION" in label: row_h = max(50, line_count * 18) 
                                        else: row_h = 30 
                                        ws.row_dimensions[current_row].height = row_h
                                        cell_a = ws.cell(row=current_row, column=1, value=f" {label}")
                                        cell_a.font, cell_a.border, cell_a.alignment = Font(size=10, bold=False), full_border, Alignment(horizontal='left', vertical='center')
                                        cell_b = ws.cell(row=current_row, column=2, value=value)
                                        cell_b.font, cell_b.border, cell_b.alignment = Font(size=10), full_border, Alignment(wrap_text=True, horizontal='left', vertical='center')
                                        current_row += 1

                                    if os.path.exists(stamp_image_file):
                                        stamp = ExcelImage(stamp_image_file)
                                        stamp.width = 110; stamp.height = 110
                                        ws.add_image(stamp, f'B{current_row + 1}') 

                                    xlsx_buf = io.BytesIO()
                                    wb.save(xlsx_buf); xlsx_buf.seek(0)
                                    zf.writestr(f"{file_count:02d}. PO {order_no}_{country_code}_{wh_code_bkg}.xlsx", xlsx_buf.read())
                                    file_count += 1
                                    
                            if file_count > 1:
                                add_log(st.session_state.current_user, f"Extract Forms: Generated {file_count-1} x Air Booking Excel forms")
                                st.session_state.air_bkg_res = {"count": file_count - 1, "zip_data": zip_buf.getvalue()}
                                
                        except Exception as e:
                            st.error(f"❌ Error during processing: {e}")

        if st.session_state.air_bkg_res:
            res = st.session_state.air_bkg_res
            st.success(f"✅ Generated {res['count']} Excel files successfully!")
            st.download_button("📥 Download Generated Excel Files (ZIP)", res["zip_data"], "Air_Booking_Forms.zip", use_container_width=True)
            
        if c_air2.button("🗑️ Clear Files", key="c_air", use_container_width=True): 
            clear_files()
            st.rerun()