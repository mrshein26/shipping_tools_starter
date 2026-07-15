import streamlit as st
import pandas as pd
import zipfile
import fitz  # PyMuPDF
import io
import os
import re
import zipfile
import pdfplumber
import subprocess
import tempfile
import shutil
from datetime import datetime, timedelta
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from pypdf import PdfReader, PdfWriter
from pdf2image import convert_from_bytes
import pytesseract
import openpyxl
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter

# 💡 ဤနေရာတွင် send_discord_alert ကို ထပ်မံပေါင်းထည့်လိုက်ပါပြီ
from core.utils import add_log, extract_files, send_discord_alert

# =========================================================
# HELPER: GHOSTSCRIPT COMPRESSOR
# =========================================================
def compress_pdf_bytes(pdf_bytes, gs_setting="/ebook"):
    """Ghostscript အသုံးပြု၍ PDF Bytes များကို Compress လုပ်ပေးသော Helper Function"""
    if not shutil.which("gs"):
        return pdf_bytes 
        
    with tempfile.TemporaryDirectory() as temp_dir:
        input_path = os.path.join(temp_dir, "input.pdf")
        output_path = os.path.join(temp_dir, "output.pdf")
        
        with open(input_path, "wb") as f:
            f.write(pdf_bytes)
            
        gs_cmd = [
            "gs", "-sDEVICE=pdfwrite", "-dCompatibilityLevel=1.4",
            f"-dPDFSETTINGS={gs_setting}", "-dNOPAUSE", "-dQUIET", "-dBATCH",
            f"-sOutputFile={output_path}", input_path
        ]
        
        try:
            subprocess.run(gs_cmd, check=True)
            if os.path.exists(output_path):
                with open(output_path, "rb") as f:
                    comp_bytes = f.read()
                    if len(comp_bytes) < len(pdf_bytes):
                        return comp_bytes
        except Exception:
            pass
            
    return pdf_bytes

def render_history_ui():
    st.subheader("🛠️ Admin Tools")
    
    # =========================================================================
    # 🗜️ GLOBAL PDF COMPRESSION SETTINGS (Expander မပါဘဲ တိုက်ရိုက်ပြသခြင်း)
    # =========================================================================
    st.markdown("""
        <style>
        /* 💡 Toggle နှင့် Selectbox ကို တစ်တန်းတည်း ညီနေစေရန် */
        div[data-testid="stHorizontalBlock"] {
            align-items: center !important;
            margin-bottom: -10px !important;
        }
        </style>
    """, unsafe_allow_html=True)

    c_comp1, c_comp2 = st.columns([1.5, 2.5], vertical_alignment="center")
    
    with c_comp1:
        do_compress = st.toggle("Compress Output PDFs", value=True)
        
    with c_comp2:
        quality_labels = {
            "📱 Medium (150 DPI - Standard)": "/ebook",
            "📉 Low (72 DPI - Smallest)": "/screen",
            "🖨️ High (300 DPI - Print)": "/printer",
            "🎨 Prepress (Max Quality)": "/prepress"
        }
        # 💡 "Quality" label ကို လုံးဝဝှက်ထားပြီး Selectbox ကိုသာ ပြသပါမည်
        selected_qual = st.selectbox("Quality", list(quality_labels.keys()), index=0, disabled=not do_compress, label_visibility="collapsed")
        gs_setting = quality_labels[selected_qual] if do_compress else None
        
    if do_compress and not shutil.which("gs"):
        st.warning("⚠️ Ghostscript ကို install မလုပ်ထားပါ။ Compression အလုပ်လုပ်မည် မဟုတ်ပါ။")

    if "up_key" not in st.session_state: st.session_state.up_key = 0
    if "air_merge_res" not in st.session_state: st.session_state.air_merge_res = None
    if "arc_res" not in st.session_state: st.session_state.arc_res = None
    if "pdf_files_res" not in st.session_state: st.session_state.pdf_files_res = None
    if "ro_data_res" not in st.session_state: st.session_state.ro_data_res = None
    if "rex_data_res" not in st.session_state: st.session_state.rex_data_res = None
    if "hm_ext_res" not in st.session_state: st.session_state.hm_ext_res = None
    if "lic_res" not in st.session_state: st.session_state.lic_res = None
    if "import_lists_res" not in st.session_state: st.session_state.import_lists_res = None
    if "export_summary_res" not in st.session_state: st.session_state.export_summary_res = None

    def clear_files(): 
        st.session_state.up_key += 1
        st.session_state.air_merge_res = None
        st.session_state.arc_res = None
        st.session_state.pdf_files_res = None
        st.session_state.ro_data_res = None
        st.session_state.rex_data_res = None
        st.session_state.hm_ext_res = None
        st.session_state.lic_res = None
        st.session_state.import_lists_res = None
        st.session_state.export_summary_res = None

    # Helper Functions
    def get_alphanumeric(text):
        if not text: return ""
        return re.sub(r'[^A-Z0-9]', '', text)

    def extract_clean_text_fitz(pdf_stream):
        text = ""
        try:
            with fitz.open(stream=pdf_stream, filetype="pdf") as doc:
                for page in doc:
                    text += page.get_text() + " "
        except: pass
        return text.upper()

    def find_in_list_basic(files, target_clean):
        if not target_clean: return None
        for f in files:
            if target_clean in get_alphanumeric(f.name.upper()):
                return f
            content = extract_clean_text_fitz(f.getvalue())
            if target_clean in get_alphanumeric(content):
                return f
        return None

    def air_process_logic(asm_ups, inv_ups, lic_ups, hawb_ups, do_compress, gs_setting):
        ext_asm = extract_files(asm_ups)
        ext_inv = extract_files(inv_ups)
        ext_lic = extract_files(lic_ups)
        ext_hawb = extract_files(hawb_ups)

        zip_buf = io.BytesIO()
        p_count = 0
        
        with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for asm_f in ext_asm:
                try:
                    asm_content = asm_f.getvalue()
                    raw_asm = extract_clean_text_fitz(asm_content)
                    clean_asm = re.sub(r'\s+', '', raw_asm)
                    
                    dec_m = re.search(r'(200\d{9})', clean_asm) 
                    if not dec_m:
                        dec_m = re.search(r'DECLARATIONNO.*?(\d{12})', clean_asm)
                    dec_no = dec_m.group(1) if dec_m else f"Combined_{asm_f.name}"
                    
                    awb_m = re.search(r'(GLC[-_]?\d+)', clean_asm)
                    awb_clean = get_alphanumeric(awb_m.group(1)) if awb_m else None
                    
                    inv_m = re.search(r'(TENG[A-Z0-9\-\/]+?20\d{2})', clean_asm)
                    inv_clean = get_alphanumeric(inv_m.group(1)) if inv_m else None
                    
                    lic_m = re.search(r'(OVSEL[A-Z0-9]+)', clean_asm)
                    lic_clean = get_alphanumeric(lic_m.group(1)) if lic_m else None

                    m_inv = find_in_list_basic(ext_inv, inv_clean)
                    m_lic = find_in_list_basic(ext_lic, lic_clean)
                    
                    pkg_count = ""
                    if m_inv:
                        inv_text = extract_clean_text_fitz(m_inv.getvalue())
                        pkg_m = re.search(r'TOTAL(?:CTNS|CTN)(\d+)|TOTAL(\d+)(?:CTNS|CTN)', re.sub(r'\s+', '', inv_text))
                        if pkg_m: pkg_count = pkg_m.group(1) or pkg_m.group(2)
                    
                    if not pkg_count:
                        pkg_m_asm = re.search(r'PACKAGES(\d+)', clean_asm)
                        if pkg_m_asm: pkg_count = pkg_m_asm.group(1)

                    out_pdf = fitz.open()
                    with fitz.open(stream=asm_content, filetype="pdf") as d: out_pdf.insert_pdf(d)
                    
                    if m_inv:
                        with fitz.open(stream=m_inv.getvalue(), filetype="pdf") as d: out_pdf.insert_pdf(d)
                    if m_lic:
                        with fitz.open(stream=m_lic.getvalue(), filetype="pdf") as d: out_pdf.insert_pdf(d)
                    
                    hawb_added = False
                    if awb_clean and ext_hawb:
                        for h_f in ext_hawb:
                            imgs = convert_from_bytes(h_f.getvalue(), dpi=150)
                            for i, img in enumerate(imgs):
                                ocr_txt = pytesseract.image_to_string(img).upper()
                                page_alpha = get_alphanumeric(ocr_txt)
                                
                                if awb_clean in page_alpha:
                                    if "BOOKINGNOTE" not in page_alpha:
                                        with fitz.open(stream=h_f.getvalue(), filetype="pdf") as d:
                                            out_pdf.insert_pdf(d, from_page=i, to_page=i)
                                        hawb_added = True

                    f_name = f"{dec_no}_{pkg_count}.pdf" if pkg_count else f"{dec_no}.pdf"
                    
                    final_pdf_bytes = out_pdf.write()
                    out_pdf.close()
                    
                    if do_compress and gs_setting:
                        final_pdf_bytes = compress_pdf_bytes(final_pdf_bytes, gs_setting)
                        
                    zf.writestr(f_name, final_pdf_bytes)
                    p_count += 1
                except Exception as e:
                    pass
        return zip_buf.getvalue(), p_count

    def extract_hm_data(text, index):
        data = {}
        data['Sr No.'] = index
        
        order_match = re.search(r"H&M Order No:\s*([0-9-]+)", text)
        data['H&M Order No'] = order_match.group(1) if order_match else ""
        
        inv_no_match = re.search(r"Number:\s*([A-Z0-9-]+)", text)
        data['Invoice No'] = inv_no_match.group(1) if inv_no_match else ""
        
        inv_date_match = re.search(r"Date:\s*([\d-]+)", text)
        data['Invoice Date'] = inv_date_match.group(1) if inv_date_match else ""
        
        dest_match = re.search(r"Final Destination:\s*(.*?)(?=\n\s*(?:The following table|Port of Loading|Warehouse ID|Container No))", text, re.IGNORECASE | re.DOTALL)
        
        if dest_match:
            raw_lines = dest_match.group(1).split('\n')
            skip_words = ["terms of delivery", "sea", "air", "fca", "cpt"]
            clean_lines = [
                line.strip() for line in raw_lines 
                if line.strip() and not any(sw in line.lower() for sw in skip_words)
            ]
            
            if len(clean_lines) >= 2:
                data['Final Destination Name'] = clean_lines[0]
                data['Final Destination Code'] = clean_lines[1]
            elif len(clean_lines) == 1:
                data['Final Destination Name'] = clean_lines[0]
                data['Final Destination Code'] = ""
        else:
            data['Final Destination Name'] = ""
            data['Final Destination Code'] = ""
            
        wh_match = re.search(r"\b([A-Z]{2,3}W\d{3})\b", text)
        data['Warehouse ID'] = wh_match.group(1) if wh_match else ""
        
        raw_desc = ""
        desc_match = re.search(r'(?:Cartons|Description of Goods)(.*?)(?:HS Code|Container No|Total)', text, re.DOTALL | re.IGNORECASE)
        
        if desc_match:
            raw_desc = desc_match.group(1)
        else:
            alt_desc_match = re.search(r'Description of Goods.*?\n(.*)', text, re.DOTALL | re.IGNORECASE)
            if alt_desc_match:
                raw_desc = alt_desc_match.group(1)

        if raw_desc:
            clean_desc = re.sub(r'[\n\r]+', ' ', raw_desc)
            clean_desc = clean_desc.replace("\\n", " ").replace('"', ' ').replace(',', ' ')
            if "the scientific" in clean_desc.lower(): clean_desc = re.split(r'(?i)the scientific', clean_desc)[0]
            if "ovis aries" in clean_desc.lower(): clean_desc = re.split(r'(?i)ovis aries', clean_desc)[0]
            clean_desc = re.sub(r'(?i)gauge within.*?rows', ' ', clean_desc)
            garbage_words = r'(?i)\b(Quantity|Price|Amount|USD|Pieces|Pcs|Cartons?|Description of Goods|Container No)\b'
            clean_desc = re.sub(garbage_words, ' ', clean_desc)
            clean_desc = re.sub(r'\b[A-Z]{3}\d{3}\b', ' ', clean_desc)
            clean_desc = re.sub(r'\b\d+(?:\.\d+)?\b(?!\s*%)', ' ', clean_desc)
            data['Description of Goods'] = " ".join(clean_desc.split()).strip()
        else:
            data['Description of Goods'] = ""
            
        cartons_match = re.search(r"(\d+)\s*Cartons", text)
        data['Cartons'] = int(cartons_match.group(1)) if cartons_match else 0
        
        pieces_match = re.search(r"\b(\d+)\s+\d+\.\d{2}\s+\d+\.\d{2}\b", text)
        if pieces_match:
            data['Pieces'] = int(pieces_match.group(1))
        else:
            alt_pieces = re.search(r"(\d+)\s*(?:Pieces|Pcs)\b", text, re.IGNORECASE)
            data['Pieces'] = int(alt_pieces.group(1)) if alt_pieces else 0
        
        net_wt_match = re.search(r"Net Weight:\s*(?:\|\s*)?([\d.]+)\s*KG", text)
        data['Net Weight (KG)'] = float(net_wt_match.group(1)) if net_wt_match else 0.0
        
        gross_wt_match = re.search(r"Gross Weight:\s*(?:\|\s*)?([\d.]+)\s*KG", text)
        data['Gross Weight (KG)'] = float(gross_wt_match.group(1)) if gross_wt_match else 0.0
        
        return data

    def process_hm_files(uploaded_files):
        extracted_records = []
        index = 1 
        for uploaded_file in uploaded_files:
            if uploaded_file.name.endswith(".zip"):
                with zipfile.ZipFile(uploaded_file, "r") as z:
                    for file_name in z.namelist():
                        if file_name.lower().endswith(".pdf") and not file_name.startswith("__MACOSX"):
                            with z.open(file_name) as f:
                                pdf_stream = io.BytesIO(f.read())
                                with pdfplumber.open(pdf_stream) as pdf:
                                    text = "\n".join([page.extract_text() for page in pdf.pages if page.extract_text()])
                                    extracted_records.append(extract_hm_data(text, index))
                                    index += 1
            elif uploaded_file.name.endswith(".pdf"):
                pdf_stream = io.BytesIO(uploaded_file.getvalue())
                with pdfplumber.open(pdf_stream) as pdf:
                     text = "\n".join([page.extract_text() for page in pdf.pages if page.extract_text()])
                     extracted_records.append(extract_hm_data(text, index))
                     index += 1
        return extracted_records

    def format_hm_excel(writer, df):
        worksheet = writer.sheets['HM_Extracted_Data']
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        border_side = Side(border_style="thin", color="D3D3D3")
        border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_align_indent = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        right_align_indent = Alignment(horizontal="right", vertical="center", wrap_text=True, indent=1)

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = border

        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.freeze_panes = "A2"

        for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column), start=2):
            worksheet.row_dimensions[row_idx].height = 17
            for cell in row:
                cell.border = border
                if cell.column == 1: 
                    cell.alignment = center_align
                elif cell.column in [3, 8]:
                    cell.alignment = left_align_indent
                elif cell.column >= 9:  
                    cell.alignment = right_align_indent
                    if cell.column >= 11: cell.number_format = '#,##0.00' 
                    else: cell.number_format = '#,##0' 
                else:
                    cell.alignment = center_align
                    
        worksheet.column_dimensions['A'].width = 6
        worksheet.column_dimensions['H'].width = 45  
        for col in worksheet.columns:
            col_letter = col[0].column_letter
            if col_letter not in ['A', 'H']: 
                worksheet.column_dimensions[col_letter].width = 20

    # =========================================================================
    # 🌟 UI Layout (Main Tabs: Import Tools, Export Tools, & User Records)
    # =========================================================================
    tab_import, tab_export, tab_records = st.tabs(["📥 Import Tools", "📤 Export Tools", "👨🏻‍💻 User Records"])

    # ---------------------------------------------------------
    # 📥 IMPORT TOOLS TAB
    # ---------------------------------------------------------
    with tab_import:
        import_tool = st.radio(
            "Select Import Tool",
            ["Import Data (RO)", "License Balance", "Extract Import Lists"],
            horizontal=True,
            label_visibility="collapsed"
        )
        if import_tool == "Import Data (RO)":
            st.info("Upload multiple **PDFs** or **ZIP files** to extract RO Import Data into Excel and Rename PDFs.")
            ro_files = st.file_uploader("Upload RO PDFs or ZIP", type=["pdf", "zip"], accept_multiple_files=True, key=f"ro_{st.session_state.up_key}")
            
            c_ro1, c_ro2 = st.columns(2)
            if c_ro1.button("⚡ Extract & Rename", use_container_width=True, key="btn_ro_extract"):
                if not ro_files: 
                    st.error("Please upload files.")
                else:
                    with st.spinner("Extracting RO Data & Renaming PDFs..."):
                        extracted = extract_files(ro_files)
                        all_data = []
                        renamed_pdfs = [] 
                        
                        def c_num(v):
                            try: return float(re.sub(r'[^\d.]', '', str(v))) if re.sub(r'[^\d.]', '', str(v)) else None
                            except: return None
                        def c_int(v):
                            try: return int(re.sub(r'[^\d]', '', str(v))) if re.sub(r'[^\d]', '', str(v)) else None
                            except: return None

                        for i, f_obj in enumerate(extracted, 1):
                            f_obj.seek(0)
                            raw_pdf_bytes = f_obj.read() 
                            
                            try:
                                doc = fitz.open(stream=raw_pdf_bytes, filetype="pdf")
                                text = "".join([p.get_text("text", sort=True) for p in doc])
                                cl_txt = text.replace('\n', ' ')

                                m_dec = re.search(r"(1003\d{8})", text)
                                m_date = re.search(r"(202\d/\d{2}/\d{2})", text)
                                m_inv = re.search(r"Invoice[\s\S]{0,100}?A[\s\n\-\:]*([A-Z0-9]+)", text, re.IGNORECASE)
                                
                                t_ctn = ""
                                m_pk = re.search(r"([\d,]+)[\s\n]*PK", text, re.IGNORECASE)
                                if m_pk: t_ctn = m_pk.group(1)
                                else:
                                    m_pkg = re.search(r"Packages[\s\S]{0,30}?(?<![\d,])([\d,]+)(?![\d,])", text, re.IGNORECASE)
                                    if m_pkg: t_ctn = m_pkg.group(1)

                                kgms = re.findall(r"([\d\.,]+)[\s\n]*KGM", text, re.IGNORECASE)
                                g_wt = kgms[0] if len(kgms) > 0 else ""
                                n_wt = kgms[1] if len(kgms) > 1 else ""

                                m_cif = re.search(r"CIF[\s\-\w]*USD[\s\-]*([\d\.,]+)", cl_txt, re.IGNORECASE)
                                if not m_cif: m_cif = re.search(r"Invoice price[\s\S]{0,50}?USD[\s\-]*([\d\.,]+)", cl_txt, re.IGNORECASE)

                                dec_str = m_dec.group(1) if m_dec else "UNKNOWN"
                                ctn_str = str(t_ctn).replace("/", "-").strip() if t_ctn else "0"
                                if "PK" not in ctn_str.upper():
                                    ctn_str += " PK"
                                
                                new_pdf_name = f"{i}_{dec_str}_{ctn_str}.pdf"
                                
                                final_pdf_bytes = raw_pdf_bytes
                                if do_compress and gs_setting:
                                    final_pdf_bytes = compress_pdf_bytes(final_pdf_bytes, gs_setting)
                                    
                                renamed_pdfs.append((new_pdf_name, final_pdf_bytes))

                                all_data.append({
                                    "No.": i, "Declaration no.": c_int(m_dec.group(1) if m_dec else ""),
                                    "Declaration date": m_date.group(1) if m_date else "", "Invoice No.": m_inv.group(1) if m_inv else "",
                                    "Description": "Yarn and Accessories", "Total Ctn": c_num(t_ctn),
                                    "Net weight": c_num(n_wt), "Gross weight": c_num(g_wt), "CIF Value (USD)": c_num(m_cif.group(1) if m_cif else "")
                                })
                            except: pass

                        if all_data:
                            cols = ["No.", "Declaration no.", "Declaration date", "Invoice No.", "Description", "Total Ctn", "Net weight", "Gross weight", "CIF Value (USD)"]
                            df = pd.DataFrame(all_data)[cols]
                            
                            tmp_buf = io.BytesIO()
                            df.to_excel(tmp_buf, index=False, startrow=1)
                            tmp_buf.seek(0)
                            
                            wb = load_workbook(tmp_buf)
                            ws = wb.active
                            ws.merge_cells('A1:I1')
                            ws['A1'] = "TENG HUI (MYANMAR) KNITTING CO.,LTD_IMPORT DATA LISTS"
                            ws['A1'].font, ws['A1'].alignment = Font(bold=True, size=14), Alignment(horizontal="center", vertical="center")

                            h_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
                            for cell in ws[2]:
                                if cell.value: cell.value = str(cell.value).upper()
                                cell.font, cell.alignment, cell.fill = Font(bold=True), Alignment(horizontal="center", vertical="center"), h_fill

                            ws.auto_filter.ref = f"A2:I{ws.max_row}"
                            ws.freeze_panes = 'A3'

                            for col in ws.columns:
                                c_let = get_column_letter(col[0].column)
                                if c_let == 'A': ws.column_dimensions[c_let].width = 6
                                elif c_let == 'B': ws.column_dimensions[c_let].width = 16
                                else: ws.column_dimensions[c_let].width = max([len(str(c.value)) for c in col] + [0]) + 2

                            l_row = ws.max_row
                            if l_row >= 3:
                                s_row = l_row + 1
                                ws[f'E{s_row}'] = "TOTAL :"
                                ws[f'E{s_row}'].font, ws[f'E{s_row}'].alignment = Font(bold=True), Alignment(horizontal="right", vertical="center")
                                for c_let in ['F', 'G', 'H', 'I']:
                                    ws[f'{c_let}{s_row}'] = f"=SUM({c_let}3:{c_let}{l_row})"
                                    ws[f'{c_let}{s_row}'].font = Font(bold=True)
                                    ws[f'{c_let}{s_row}'].number_format = '#,##0.00' if c_let in ['G', 'H', 'I'] else '#,##0'

                            out_xl = io.BytesIO()
                            wb.save(out_xl)
                            out_xl.seek(0)

                            out_zip = io.BytesIO()
                            with zipfile.ZipFile(out_zip, "w", zipfile.ZIP_DEFLATED) as zf:
                                zf.writestr("Teng_Hui_RO_Import_Data.xlsx", out_xl.getvalue())
                                for pdf_name, pdf_bytes in renamed_pdfs:
                                    zf.writestr(pdf_name, pdf_bytes)
                            
                            out_zip.seek(0)

                            # 💡 Discord Alert
                            mm_time = datetime.utcnow() + timedelta(hours=6, minutes=30)
                            try:
                                send_discord_alert(st.session_state.current_user, "📊 **RO Import Data:** Successfully extracted RO Data & Renamed PDFs.", mm_time.strftime("%Y-%m-%d %H:%M:%S"))
                            except: pass
                            
                            st.session_state.ro_data_res = out_zip.getvalue()
                        else: st.warning("No data extracted.")

            if st.session_state.ro_data_res:
                st.success("✅ RO Data Excel & Renamed PDFs successfully generated!")
                st.download_button("📥 Download Excel & Renamed PDFs (ZIP)", st.session_state.ro_data_res, "Renamed_PDFs_and_RO_Data.zip", mime="application/zip", use_container_width=True)
                
            if c_ro2.button("🗑️ Clear Files", key="c_ro_btn", use_container_width=True): clear_files(); st.rerun()

        elif import_tool == "License Balance":
            st.info("Upload **License PDF(s) or ZIP** to auto-extract data and generate calculated **Balance Sheet Excel**.")
            
            lic_files = st.file_uploader("Upload License PDFs or ZIP", type=["pdf", "zip"], accept_multiple_files=True, key=f"lic_{st.session_state.up_key}")
            
            c_lic1, c_lic2 = st.columns(2)
            if c_lic1.button("⚡ Generate Excel", use_container_width=True, key="btn_lic_generate"):
                if not lic_files:
                    st.error("❌ ကျေးဇူးပြု၍ License PDF ဖိုင်များကို Upload တင်ပေးပါ။")
                else:
                    with st.spinner("Extracting PDF data & Generating Balance Sheet..."):
                        def extract_license_data(pdf_obj):
                            with pdfplumber.open(pdf_obj) as pdf:
                                full_text = "\n".join([page.extract_text() or "" for page in pdf.pages])

                            lic_no_match = re.search(r'(OVSIL\d+)', full_text)
                            lic_no = lic_no_match.group(1) if lic_no_match else "UNKNOWN"

                            issue_match = re.search(r'Start valid date\s*(\d{4}/\d{2}/\d{2})', full_text)
                            issue_date = issue_match.group(1) if issue_match else ""

                            exp_match = re.search(r'Last valid date\s*(\d{4}/\d{2}/\d{2})', full_text)
                            exp_date = exp_match.group(1) if exp_match else ""

                            val_match = re.search(r'Total amount[\s\S]{1,50}?([\d,]+\.\d+)\s*USD', full_text)
                            total_val = float(val_match.group(1).replace(',','')) if val_match else 0.0

                            items = []
                            blocks = re.split(r'Item No\.\s*\d+\s*Hs code', full_text)[1:]
                            
                            for i, block in enumerate(blocks, 1):
                                hs_match = re.search(r'^\s*(\d+)', block)
                                hs_code = hs_match.group(1) if hs_match else ""

                                desc_match = re.search(r'Description of goods\s*\n(.*?)(?=\n(?:Unit price|The following table|Quantity))', block, re.DOTALL)
                                desc = desc_match.group(1).replace('\n', ' ').strip() if desc_match else ""

                                u_price = 0.0
                                u_match = re.search(r'(?:Unit price|").*?(\d+\.\d{3})\b', block, re.DOTALL)
                                if u_match: u_price = float(u_match.group(1).replace(',', ''))

                                qty = 0.0
                                q_match = re.search(r'(?:Quantity|").*?(\d{1,3}(?:,\d{3})*\.\d{2})\b', block, re.DOTALL)
                                if q_match: qty = float(q_match.group(1).replace(',', ''))

                                used_qty = 0.0
                                used_match = re.search(r'Used total quantity.*?(\d{1,3}(?:,\d{3})*\.\d{2})\b', block, re.DOTALL | re.IGNORECASE)
                                if used_match: 
                                    used_qty = float(used_match.group(1).replace(',', ''))
                                else:
                                    zero_match = re.search(r'Used total quantity.*?0\.00\b', block, re.DOTALL | re.IGNORECASE)
                                    if zero_match: used_qty = 0.0

                                items.append((i, hs_code, desc, u_price, qty, used_qty))

                            return {
                                "License No": lic_no,
                                "Issue Date": issue_date,
                                "Expiry Date": exp_date,
                                "Total License Value": total_val,
                                "Items": items
                            }

                        generated_files = []

                        for f_obj in lic_files:
                            f_obj.seek(0)
                            try:
                                data = extract_license_data(f_obj)
                                
                                wb = openpyxl.Workbook()
                                ws = wb.active
                                ws.title = "Balance Sheet"
                                ws.views.sheetView[0].showGridLines = True

                                HEADER_FILL = PatternFill(start_color="1F385C", end_color="1F385C", fill_type="solid")
                                ZEBRA_FILL = PatternFill(start_color="F4F7FA", end_color="F4F7FA", fill_type="solid")
                                TOTAL_FILL = PatternFill(start_color="E6EDF5", end_color="E6EDF5", fill_type="solid")
                                FONT_TITLE = Font(name="Segoe UI", size=16, bold=True, color="1F385C")
                                FONT_SUBTITLE = Font(name="Segoe UI", size=10, italic=True, color="555555")
                                FONT_HEADER = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
                                FONT_BODY = Font(name="Segoe UI", size=10)
                                FONT_BODY_BOLD = Font(name="Segoe UI", size=10, bold=True)
                                FONT_TOTAL = Font(name="Segoe UI", size=11, bold=True, color="1F385C")
                                ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
                                ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
                                ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
                                THIN_BORDER_SIDE = Side(border_style="thin", color="D0D7DE")
                                THIN_BORDER = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE)
                                HEADER_BORDER = Border(left=Side(border_style="thin", color="3A5A80"), right=Side(border_style="thin", color="3A5A80"))
                                TOTAL_BORDER = Border(top=Side(border_style="thin", color="1F385C"), bottom=Side(border_style="double", color="1F385C"))

                                ws['A1'] = "TENG HUI (MYANMAR) KNITTING COMPANY LIMITED"
                                ws['A1'].font = FONT_TITLE
                                ws['A2'] = f"Import License Balance Sheet | License No: {data['License No']}"
                                ws['A2'].font = FONT_SUBTITLE

                                ws['A4'] = "License Details"
                                ws['A4'].font = Font(name="Segoe UI", size=12, bold=True, color="1F385C")

                                details = [
                                    ("License No:", data['License No']),
                                    ("Issue Date:", data['Issue Date']),
                                    ("Expiry Date:", data['Expiry Date']),
                                    ("Total License Value:", data['Total License Value'])
                                ]

                                for i, (lbl, val) in enumerate(details):
                                    row = 5 + i
                                    ws[f'A{row}'] = lbl
                                    ws[f'A{row}'].font = FONT_BODY_BOLD
                                    ws[f'B{row}'] = val
                                    ws[f'B{row}'].font = FONT_BODY
                                    if isinstance(val, float):
                                        ws[f'B{row}'].number_format = '$#,##0.00'
                                    ws[f'B{row}'].alignment = ALIGN_LEFT

                                headers = [
                                    "Item No.", "HS Code", "Description of Goods", "Unit Price (USD)",
                                    "Original Quantity", "Used Total Quantity", "Balance Quantity", "Balance Amount (USD)"
                                ]

                                start_row = 11
                                for col_idx, text in enumerate(headers, start=1):
                                    cell = ws.cell(row=start_row, column=col_idx, value=text)
                                    cell.font = FONT_HEADER
                                    cell.fill = HEADER_FILL
                                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                                    cell.border = HEADER_BORDER

                                ws.row_dimensions[start_row].height = 28

                                current_row = start_row + 1
                                for item in data['Items']:
                                    item_no, hs_code, desc, unit_price, qty, used_qty = item
                                    
                                    ws.cell(row=current_row, column=1, value=item_no).alignment = ALIGN_CENTER
                                    ws.cell(row=current_row, column=2, value=hs_code).alignment = ALIGN_CENTER
                                    ws.cell(row=current_row, column=3, value=desc).alignment = ALIGN_LEFT
                                    
                                    cell_price = ws.cell(row=current_row, column=4, value=unit_price)
                                    cell_price.number_format = '#,##0.000'
                                    cell_price.alignment = ALIGN_RIGHT
                                    
                                    cell_qty = ws.cell(row=current_row, column=5, value=qty)
                                    cell_qty.number_format = '#,##0.00'
                                    cell_qty.alignment = ALIGN_RIGHT
                                    
                                    cell_used = ws.cell(row=current_row, column=6, value=used_qty)
                                    cell_used.number_format = '#,##0.00'
                                    cell_used.alignment = ALIGN_RIGHT
                                    
                                    cell_bal_qty = ws.cell(row=current_row, column=7, value=f"=E{current_row}-F{current_row}")
                                    cell_bal_qty.number_format = '#,##0.00'
                                    cell_bal_qty.alignment = ALIGN_RIGHT
                                    
                                    cell_bal_amt = ws.cell(row=current_row, column=8, value=f"=G{current_row}*D{current_row}")
                                    cell_bal_amt.number_format = '$#,##0.00'
                                    cell_bal_amt.alignment = ALIGN_RIGHT
                                    
                                    for col_idx in range(1, 9):
                                        c = ws.cell(row=current_row, column=col_idx)
                                        c.font = FONT_BODY
                                        c.border = THIN_BORDER
                                        if current_row % 2 == 1:
                                            c.fill = ZEBRA_FILL
                                            
                                    ws.row_dimensions[current_row].height = 20
                                    current_row += 1

                                ws.cell(row=current_row, column=1, value="Total").font = FONT_TOTAL
                                ws.cell(row=current_row, column=1).alignment = ALIGN_LEFT

                                ws.cell(row=current_row, column=5, value=f"=SUM(E12:E{current_row-1})").number_format = '#,##0.00'
                                ws.cell(row=current_row, column=6, value=f"=SUM(F12:F{current_row-1})").number_format = '#,##0.00'
                                ws.cell(row=current_row, column=7, value=f"=SUM(G12:G{current_row-1})").number_format = '#,##0.00'
                                ws.cell(row=current_row, column=8, value=f"=SUM(H12:H{current_row-1})").number_format = '$#,##0.00'

                                for col_idx in range(1, 9):
                                    c = ws.cell(row=current_row, column=col_idx)
                                    if col_idx not in [2, 3, 4]:
                                        c.font = FONT_TOTAL
                                        if col_idx >= 5: c.alignment = ALIGN_RIGHT
                                    c.fill = TOTAL_FILL
                                    c.border = TOTAL_BORDER

                                ws.row_dimensions[current_row].height = 24
                                ws.freeze_panes = "A12"

                                for col in ws.columns:
                                    max_len = 0
                                    col_letter = get_column_letter(col[0].column)
                                    for cell in col:
                                        if cell.row < 11: continue
                                        if cell.value:
                                            max_len = max(max_len, len(str(cell.value)))
                                    if col_letter == 'C': ws.column_dimensions[col_letter].width = 45
                                    elif col_letter in ['E', 'F', 'G', 'H']: ws.column_dimensions[col_letter].width = 20
                                    else: ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

                                out_xl = io.BytesIO()
                                wb.save(out_xl)
                                out_xl.seek(0)
                                
                                safe_exp_date = data['Expiry Date'].replace('/', '-') if data['Expiry Date'] else "Unknown_Exp"
                                excel_filename = f"{data['License No']}_Exp_{safe_exp_date}.xlsx"
                                
                                generated_files.append((excel_filename, out_xl.getvalue()))

                            except Exception as e:
                                pass

                        if len(generated_files) == 1:
                            st.session_state.lic_res = {"type": "single", "name": generated_files[0][0], "data": generated_files[0][1]}
                        elif len(generated_files) > 1:
                            zip_buf = io.BytesIO()
                            with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
                                for name, file_data in generated_files:
                                    zf.writestr(name, file_data)
                            st.session_state.lic_res = {"type": "zip", "name": "License_Balance_Sheets.zip", "data": zip_buf.getvalue()}

                        # 💡 Discord Alert
                        mm_time = datetime.utcnow() + timedelta(hours=6, minutes=30)
                        try:
                            send_discord_alert(st.session_state.current_user, f"⚖️ **License Balance:** Successfully generated {len(generated_files)} Balance Sheet(s).", mm_time.strftime("%Y-%m-%d %H:%M:%S"))
                        except: pass

            if st.session_state.lic_res:
                res = st.session_state.lic_res
                st.success("✅ Balance Sheet Excel generated successfully!")
                st.download_button(
                    label="📥 Download Extracted Balance Sheet" if res["type"] == "single" else "📥 Download All Balance Sheets (ZIP)",
                    data=res["data"],
                    file_name=res["name"],
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if res["type"] == "single" else "application/zip",
                    key="btn_lic_download",
                    use_container_width=True
                )
                
            if c_lic2.button("🗑️ Clear Files", key="c_lic_clear", use_container_width=True): clear_files(); st.rerun()

        elif import_tool == "Extract Import Lists":
            import zipfile # ZIP ဖိုင်ပြုလုပ်ရန်
            
            st.info("Upload multiple **Packing List (Excel)** files to generate a consolidated summary grouped by Material Composition.")
            
            combine_files = st.toggle("ဖိုင်အားလုံးကို ပေါင်း၍ တစ်စောင်တည်းထုတ်မည် (Combine Output)", value=True)
            
            import_files = st.file_uploader("Upload Packing Lists (Excel)", type=["xlsx"], accept_multiple_files=True, key=f"import_lists_{st.session_state.up_key}")
            
            c_imp1, c_imp2 = st.columns(2)
            if c_imp1.button("⚡ Extract & Group", use_container_width=True, key="btn_imp_extract"):
                if not import_files: 
                    st.error("❌ ကျေးဇူးပြု၍ Packing List Excel ဖိုင်(များ)ကို Upload တင်ပေးပါ။")
                else:
                    with st.spinner("Extracting and processing data..."):
                        try:
                            start_row = 14 
                            
                            def generate_excel_from_rows(rows, inv_no):
                                df = pd.DataFrame(rows)
                                summary_df = df.groupby("Description", sort=False).sum().reset_index()
                                summary_df.insert(0, 'Sr. No.', range(1, 1 + len(summary_df)))
                                
                                totals = {"Sr. No.": "", "Description": "TOTAL:"}
                                for col in summary_df.columns[2:]: totals[col] = summary_df[col].sum()
                                summary_df.loc[len(summary_df)] = totals
                                
                                wb = openpyxl.Workbook()
                                ws = wb.active
                                ws.title = "Consolidated Summary" if combine_files else f"{inv_no} Summary"
                                
                                headers = list(summary_df.columns)
                                ws.append(headers)
                                
                                header_font = Font(bold=True, color="FFFFFF")
                                header_fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
                                thin_border = Border(left=Side(style='thin', color="DDDDDD"), right=Side(style='thin', color="DDDDDD"), top=Side(style='thin', color="DDDDDD"), bottom=Side(style='thin', color="DDDDDD"))
                                
                                for col_num, header in enumerate(headers, 1):
                                    cell = ws.cell(row=1, column=col_num)
                                    cell.font, cell.fill, cell.alignment, cell.border = header_font, header_fill, Alignment(horizontal="center", vertical="center"), thin_border
                                
                                for r_idx, row in enumerate(dataframe_to_rows(summary_df, index=False, header=False), 2):
                                    for c_idx, value in enumerate(row, 1):
                                        cell = ws.cell(row=r_idx, column=c_idx, value=value)
                                        cell.border = thin_border
                                        if c_idx == 1: cell.alignment = Alignment(horizontal="center", vertical="center")
                                        elif c_idx == 2:
                                            cell.alignment = Alignment(horizontal="left", vertical="center")
                                            if "TOTAL:" in str(value): cell.alignment = Alignment(horizontal="right", vertical="center")
                                        else:
                                            cell.alignment = Alignment(horizontal="right", vertical="center")
                                            if isinstance(value, (int, float)) and value != "":
                                                if "Volume" in headers[c_idx-1]: cell.number_format = '0.000'
                                                elif "Carton(S)" in headers[c_idx-1] or "Sr. No." in headers[c_idx-1]: cell.number_format = '0'
                                                else: cell.number_format = '0.00'
                                
                                total_row_idx = len(summary_df) + 1
                                for c_idx in range(1, len(headers) + 1):
                                    cell = ws.cell(row=total_row_idx, column=c_idx)
                                    cell.font, cell.fill = Font(bold=True), PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
                                
                                ws.column_dimensions['A'].width = 8; ws.column_dimensions['B'].width = 45; ws.column_dimensions['C'].width = 12
                                ws.column_dimensions['D'].width = 15; ws.column_dimensions['E'].width = 15; ws.column_dimensions['F'].width = 15; ws.column_dimensions['G'].width = 15
                                
                                out_xl = io.BytesIO(); wb.save(out_xl); out_xl.seek(0)
                                return out_xl.getvalue()

                            # ==========================================
                            # (၁) ခလုတ် ဖွင့်ထားလျှင် (ပေါင်းမည်)
                            # ==========================================
                            if combine_files:
                                combined_rows = []
                                combined_invoice_no = "Combined_Invoice"
                                
                                for import_file in import_files:
                                    df_raw = pd.read_excel(import_file, header=None, engine='openpyxl')
                                    inv_no = "Invoice"
                                    
                                    for r in range(min(12, len(df_raw))):
                                        for c in range(len(df_raw.columns)):
                                            cell_val = str(df_raw.iloc[r, c]).strip()
                                            if "INVOICE NO" in cell_val.upper():
                                                for next_c in range(c + 1, min(c + 4, len(df_raw.columns))):
                                                    val = str(df_raw.iloc[r, next_c]).strip()
                                                    if val and val != "nan":
                                                        inv_no = val
                                                        if combined_invoice_no == "Combined_Invoice": combined_invoice_no = val
                                                        break
                                                break
                                                
                                    for i in range(start_row, len(df_raw)):
                                        desc_val = str(df_raw.iloc[i, 0]).strip()
                                        if desc_val.upper() == 'TOTAL:': break
                                        
                                        cartons = df_raw.iloc[i, 4]
                                        if pd.notna(cartons) and str(cartons).replace('.', '', 1).isdigit():
                                            composition = desc_val
                                            if i + 1 < len(df_raw) and pd.isna(df_raw.iloc[i+1, 4]) and pd.notna(df_raw.iloc[i+1, 0]):
                                                 comp = str(df_raw.iloc[i+1, 0]).strip()
                                                 if comp and comp != 'nan': composition = comp
                                                 
                                            inv_upper, desc_upper = str(inv_no).upper(), composition.upper()
                                            if "%" in composition:
                                                if inv_upper.startswith("LY") or "CUT PIECES" in desc_upper:
                                                    composition = f"{composition} KNITTING PANEL"
                                                elif inv_upper.startswith("PX") or "YARN" in desc_upper:
                                                    composition = f"{composition} KNITTING YARN"
                                                 
                                            combined_rows.append({
                                                "Description": composition, "Carton(S)": float(df_raw.iloc[i, 4]),
                                                "Quantity": float(df_raw.iloc[i, 5]) if pd.notna(df_raw.iloc[i, 5]) else 0.0,
                                                "N.W (KGs)": float(df_raw.iloc[i, 6]) if pd.notna(df_raw.iloc[i, 6]) else 0.0,
                                                "G.W (KGs)": float(df_raw.iloc[i, 7]) if pd.notna(df_raw.iloc[i, 7]) else 0.0,
                                                "Volume (CBM)": float(df_raw.iloc[i, 8]) if pd.notna(df_raw.iloc[i, 8]) else 0.0
                                            })
                                            
                                if not combined_rows: 
                                    st.warning("⚠️ ဖိုင်များထဲတွင် အချက်အလက်များ ရှာမတွေ့ပါ။")
                                else:
                                    excel_data = generate_excel_from_rows(combined_rows, combined_invoice_no)
                                    st.session_state.import_lists_res = {"mode": "combine", "name": f"{combined_invoice_no} Summary Data.xlsx", "data": excel_data, "count": len(import_files)}

                            # ==========================================
                            # (၂) ခလုတ် ပိတ်ထားလျှင် (တစ်စောင်စီကို ZIP ဖြင့်ထုတ်မည်)
                            # ==========================================
                            else:
                                processed_list = []
                                
                                for import_file in import_files:
                                    file_rows = []
                                    inv_no = "Invoice"
                                    df_raw = pd.read_excel(import_file, header=None, engine='openpyxl')
                                    
                                    for r in range(min(12, len(df_raw))):
                                        for c in range(len(df_raw.columns)):
                                            cell_val = str(df_raw.iloc[r, c]).strip()
                                            if "INVOICE NO" in cell_val.upper():
                                                for next_c in range(c + 1, min(c + 4, len(df_raw.columns))):
                                                    val = str(df_raw.iloc[r, next_c]).strip()
                                                    if val and val != "nan":
                                                        inv_no = val
                                                        break
                                                break
                                                
                                    for i in range(start_row, len(df_raw)):
                                        desc_val = str(df_raw.iloc[i, 0]).strip()
                                        if desc_val.upper() == 'TOTAL:': break
                                        
                                        cartons = df_raw.iloc[i, 4]
                                        if pd.notna(cartons) and str(cartons).replace('.', '', 1).isdigit():
                                            composition = desc_val
                                            if i + 1 < len(df_raw) and pd.isna(df_raw.iloc[i+1, 4]) and pd.notna(df_raw.iloc[i+1, 0]):
                                                 comp = str(df_raw.iloc[i+1, 0]).strip()
                                                 if comp and comp != 'nan': composition = comp
                                                 
                                            inv_upper, desc_upper = str(inv_no).upper(), composition.upper()
                                            if "%" in composition:
                                                if inv_upper.startswith("LY") or "CUT PIECES" in desc_upper:
                                                    composition = f"{composition} KNITTING PANEL"
                                                elif inv_upper.startswith("PX") or "YARN" in desc_upper:
                                                    composition = f"{composition} KNITTING YARN"
                                                 
                                            file_rows.append({
                                                "Description": composition, "Carton(S)": float(df_raw.iloc[i, 4]),
                                                "Quantity": float(df_raw.iloc[i, 5]) if pd.notna(df_raw.iloc[i, 5]) else 0.0,
                                                "N.W (KGs)": float(df_raw.iloc[i, 6]) if pd.notna(df_raw.iloc[i, 6]) else 0.0,
                                                "G.W (KGs)": float(df_raw.iloc[i, 7]) if pd.notna(df_raw.iloc[i, 7]) else 0.0,
                                                "Volume (CBM)": float(df_raw.iloc[i, 8]) if pd.notna(df_raw.iloc[i, 8]) else 0.0
                                            })
                                            
                                    if file_rows:
                                        excel_data = generate_excel_from_rows(file_rows, inv_no)
                                        processed_list.append({"name": f"{inv_no} Summary Data.xlsx", "data": excel_data})
                                        
                                if not processed_list: 
                                    st.warning("⚠️ ဖိုင်များထဲတွင် အချက်အလက်များ ရှာမတွေ့ပါ။")
                                else:
                                    # 💡 သီးသန့်ဖိုင်များကို ZIP ဖိုင် တစ်ခုတည်းအဖြစ် ပေါင်းထုပ်ခြင်း
                                    zip_buffer = io.BytesIO()
                                    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
                                        for file_info in processed_list:
                                            zip_file.writestr(file_info["name"], file_info["data"])
                                            
                                    st.session_state.import_lists_res = {
                                        "mode": "separate", 
                                        "name": "Separate_Summaries.zip", 
                                        "data": zip_buffer.getvalue(), 
                                        "count": len(import_files)
                                    }

                            # 💡 Discord Alert
                            mm_time = datetime.utcnow() + timedelta(hours=6, minutes=30)
                            try:
                                send_discord_alert(st.session_state.current_user, f"📦 **Extract Import Lists:** Successfully processed {len(import_files)} files (Combine: {combine_files}).", mm_time.strftime("%Y-%m-%d %H:%M:%S"))
                            except: pass
                                
                        except Exception as e:
                            st.error(f"❌ Error processing files: {e}")
        
            # ==========================================
            # UI တွင် Download ခလုတ်များ ဖော်ပြခြင်း (တစ်ခါတည်းသာ)
            # ==========================================
            if st.session_state.get("import_lists_res"):
                res = st.session_state.import_lists_res
                
                # ပေါင်းထားလျှင် (Excel တစ်စောင်တည်း)
                if res["mode"] == "combine":
                    st.success(f"✅ Successfully consolidated {res['count']} Excel files into one summary!")
                    st.download_button("📥 Download Consolidated Summary Excel", res["data"], res["name"], mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="btn_imp_download_combine_final", use_container_width=True)
                
                # ခွဲထုတ်ထားလျှင် (ZIP ဖိုင် တစ်ခုတည်း)
                elif res["mode"] == "separate":
                    st.success(f"✅ Successfully processed {res['count']} Excel files separately into a ZIP file!")
                    st.download_button("📥 Download All Files (ZIP)", res["data"], res["name"], mime="application/zip", key="btn_imp_download_zip_final", use_container_width=True)
                
            if c_imp2.button("🗑️ Clear Files", key="c_imp_clear_final", use_container_width=True): 
                clear_files()
                st.session_state.pop("import_lists_res", None)
                st.rerun()

    # ---------------------------------------------------------
    # 📤 EXPORT TOOLS TAB
    # ---------------------------------------------------------
    with tab_export:
        export_tool = st.radio(
            "Select Export Tool",
            ["Air Documents Merge", "Remove & Combine", "Export Data", "Rex Data", "Extract Export Summary", "Extract Inv no"],
            horizontal=True,
            label_visibility="collapsed"
        )

        if export_tool == "Air Documents Merge":
            st.info("⚠️ HAWB Scan ဖိုင်များမှ AWB No ကို OCR ဖြင့် ဖတ်၍ ပေါင်းစည်းပေးပါသည်။")
            
            c1, c2 = st.columns(2)
            with c1:
                u_asm = st.file_uploader("Upload ASM", type=["pdf", "zip"], accept_multiple_files=True, key=f"air_asm_{st.session_state.up_key}")
                u_inv = st.file_uploader("Upload INV", type=["pdf", "zip"], accept_multiple_files=True, key=f"air_inv_{st.session_state.up_key}")
            with c2:
                u_lic = st.file_uploader("Upload LIC", type=["pdf", "zip"], accept_multiple_files=True, key=f"air_lic_{st.session_state.up_key}")
                u_hawb = st.file_uploader("Upload Scanned HAWBs", type=["pdf", "zip"], accept_multiple_files=True, key=f"air_hawb_{st.session_state.up_key}")

            btn_col1, btn_col2 = st.columns(2)

            with btn_col1:
                if st.session_state.get('air_merge_res') and st.session_state.air_merge_res.get("cnt", 0) > 0:
                    res = st.session_state.air_merge_res
                    st.download_button("📥 Download Result ZIP", data=res["data"], file_name="Air_Merge_Doc.zip", mime="application/zip", use_container_width=True, key="download_air_merge_success")
                    process_clicked = False 
                else:
                    process_clicked = st.button("🔗 Process Air Docs Merge", use_container_width=True, key="process_air_docs_btn")

            with btn_col2:
                if st.button("🗑️ Clear Files", key="clear_air_merge_final", use_container_width=True):
                    st.session_state.air_merge_res = None
                    clear_files(); st.rerun()
            
            if process_clicked:
                if not u_asm: st.error("ASM ဖိုင်တင်ပေးရန် လိုအပ်ပါသည်။")
                else:
                    with st.status("⏳ OCR စနစ်ဖြင့် ဖတ်၍ ပေါင်းစည်းနေပါသည်...", expanded=True) as status:
                        res_data, res_cnt = air_process_logic(u_asm, u_inv, u_lic, u_hawb, do_compress, gs_setting)
                        if res_cnt > 0:
                            status.update(label="✅ လုပ်ဆောင်မှု ပြီးစီးပါပြီ", state="complete", expanded=False)
                            st.session_state.air_merge_res = {"data": res_data, "cnt": res_cnt}
                            
                            # 💡 Discord Alert
                            mm_time = datetime.utcnow() + timedelta(hours=6, minutes=30)
                            try:
                                send_discord_alert(st.session_state.current_user, f"✈️ **Air Docs Merge:** Merged **{res_cnt}** files.", mm_time.strftime("%Y-%m-%d %H:%M:%S"))
                            except: pass

                            st.success(f"စုစုပေါင်း {res_cnt} ဖိုင် ပေါင်းစည်းပြီးပါပြီ။")
                            st.rerun()
                        else:
                            status.update(label="⚠️ ပေါင်းစည်းရန် ဖိုင်မတွေ့ပါ", state="error", expanded=False)
                            st.error("ပေါင်းစည်းရန် ဖိုင်မတွေ့ပါ။")

        elif export_tool == "Remove & Combine":
            st.info("Upload multiple **PDFs** or **ZIP files** to remove unnecessary pages and combine them.")
            arc_files = st.file_uploader("Upload PDFs or ZIP", type=["pdf", "zip"], accept_multiple_files=True, key=f"arc_{st.session_state.up_key}")
            
            c_arc1, c_arc2 = st.columns(2)
            if c_arc1.button("⚡ Process & Combine", use_container_width=True, key="btn_arc_process"):
                if not arc_files: st.error("Please upload PDF or ZIP files.")
                else:
                    with st.spinner("Processing documents..."):
                        extracted = extract_files(arc_files)
                        if not extracted: st.error("No PDFs found.")
                        else:
                            pdf_list = []
                            for f_obj in extracted:
                                f_obj.seek(0)
                                metadata = {'filename': f_obj.name, 'file_obj': f_obj, 'export_no': None, 'elns_val': None, 'is_notification': False}
                                elns_name = re.search(r'(OVSEL[A-Z0-9\-/]+)', metadata['filename'], re.IGNORECASE)
                                if elns_name: metadata['elns_val'] = elns_name.group(1).upper()
                                
                                try:
                                    with pdfplumber.open(f_obj) as pdf:
                                        full_text = "".join([page.extract_text() + "\n" for page in pdf.pages if page.extract_text()])
                                        if full_text:
                                            export_match = re.search(r'(?:Export control No\.|Export control No|Control No).*?(\d{10,15})', full_text, re.DOTALL | re.IGNORECASE)
                                            if export_match: metadata['export_no'] = export_match.group(1)
                                            else:
                                                name_match = re.search(r'(\d{10,15})', metadata['filename'])
                                                if name_match: metadata['export_no'] = name_match.group(1)
                                                
                                            if not metadata.get('elns_val'):
                                                elns_match = re.search(r'ELNS[\s\n:-]*([A-Z0-9\-/]{5,})', full_text, re.IGNORECASE)
                                                if elns_match: metadata['elns_val'] = elns_match.group(1).strip()
                                            if not metadata.get('elns_val'):
                                                ovsel_match = re.search(r'(OVSEL[A-Z0-9\-/]+)', full_text, re.IGNORECASE)
                                                if ovsel_match: metadata['elns_val'] = ovsel_match.group(1).strip()
                                                
                                            metadata['is_notification'] = "Allowed shipment notification" in full_text
                                except: pass
                                pdf_list.append(metadata)

                            groups = {}; elns_to_eno = {}
                            for data in pdf_list:
                                eno, elns = data.get('export_no'), data.get('elns_val')
                                if eno:
                                    if eno not in groups: groups[eno] = []
                                    groups[eno].append(data)
                                    if elns: elns_to_eno[elns] = eno

                            for data in pdf_list:
                                eno, elns = data.get('export_no'), data.get('elns_val')
                                if not eno and elns and elns in elns_to_eno:
                                    target_eno = elns_to_eno[elns]
                                    if data not in groups[target_eno]: groups[target_eno].append(data)

                            zip_buf = io.BytesIO(); count = 0
                            with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
                                for eno, members in groups.items():
                                    if len(members) >= 2:
                                        writer = PdfWriter()
                                        original_notification_name = ""
                                        members.sort(key=lambda x: x.get('is_notification', False), reverse=True)
                                        found_valid_pages = False
                                        group_elns = next((m.get('elns_val') for m in members if m.get('elns_val')), "")

                                        for member in members:
                                            if member.get('is_notification') and not original_notification_name:
                                                original_notification_name = member['filename']
                                            
                                            is_pure_license_file = bool(not member.get('export_no') and member.get('elns_val')) or bool(re.search(r'(OVSEL|LICEN)', member['filename'], re.IGNORECASE))
                                            
                                            member['file_obj'].seek(0)
                                            reader = PdfReader(member['file_obj'])
                                            total_pages = len(reader.pages)
                                            
                                            for page_num, page in enumerate(reader.pages, start=1):
                                                if is_pure_license_file:
                                                    writer.add_page(page); found_valid_pages = True; continue

                                                txt = page.extract_text() or ""
                                                clean_txt = re.sub(r'[\s\-_:,]', '', txt.upper())
                                                clean_elns = re.sub(r'[\s\-_:,]', '', group_elns.upper()) if group_elns else ""
                                                
                                                targets = ["INWARDPROCESSING", "COMMERCIALINVOICE", "PACKINGLIST", "SALECONTRACT", "EXPORTLICENSE", "EXPORTLICENCE", "DEPARTMENTOFTRADE", "MINISTRYOFCOMMERCE", "OVERSEASEXPORT", "APPROVAL", "ELNS"]
                                                exclusions = ["SHIPPINGINSTRUCTION", "BOOKINGCONFIRMATION", "BOOKINGNOTE"]
                                                
                                                has_target = any(t in clean_txt for t in targets)
                                                has_elns = bool(clean_elns and len(clean_elns) > 5 and clean_elns in clean_txt)
                                                has_exclusion = any(e in clean_txt for e in exclusions)
                                                
                                                is_strong_target = has_elns or "EXPORTLICENCE" in clean_txt or "EXPORTLICENSE" in clean_txt
                                                                                                    
                                                if is_strong_target: 
                                                    is_adding = True 
                                                    has_exclusion = False 
                                                elif has_target: 
                                                    is_adding = True
                                                elif has_exclusion: 
                                                    is_adding = False
                                                    
                                                is_image = len(clean_txt) < 20
                                                if is_adding and not has_exclusion:
                                                    writer.add_page(page); found_valid_pages = True
                                                elif is_image and page_num >= total_pages - 2:
                                                    writer.add_page(page); found_valid_pages = True
                                                
                                        if found_valid_pages:
                                            save_name = original_notification_name if original_notification_name else f"{eno}_combined.pdf"
                                            out_pdf = io.BytesIO(); writer.write(out_pdf)
                                            
                                            final_pdf_bytes = out_pdf.getvalue()
                                            if do_compress and gs_setting:
                                                final_pdf_bytes = compress_pdf_bytes(final_pdf_bytes, gs_setting)
                                                
                                            zf.writestr(save_name, final_pdf_bytes); count += 1

                            if count > 0:
                                # 💡 Discord Alert
                                mm_time = datetime.utcnow() + timedelta(hours=6, minutes=30)
                                try:
                                    send_discord_alert(st.session_state.current_user, f"✂️ **ARC:** Successfully combined **{count}** PDFs.", mm_time.strftime("%Y-%m-%d %H:%M:%S"))
                                except: pass
                                
                                st.session_state.arc_res = {"count": count, "zip_data": zip_buf.getvalue()}
                            else:
                                st.warning("⚠️ No valid PDF pairs found to combine.")

            if st.session_state.arc_res:
                res = st.session_state.arc_res
                st.success(f"✅ Successfully combined {res['count']} files!")
                st.download_button("📥 Download Combined PDFs (ZIP)", res["zip_data"], "ASMs_Combined_Sea.zip", use_container_width=True)
            if c_arc2.button("🗑️ Clear Files", key="c_arc_btn", use_container_width=True): clear_files(); st.rerun()

        elif export_tool == "Export Data":
            st.info("Upload multiple **PDFs** or **ZIP files** to extract Export Data into Excel.")
            exp_files = st.file_uploader("Upload Export PDFs or ZIP", type=["pdf", "zip"], accept_multiple_files=True, key=f"exp_{st.session_state.up_key}")
            
            c_pdf1, c_pdf2 = st.columns(2)
            if c_pdf1.button("⚡ Extract & Rename", use_container_width=True, key="btn_pdf_extract"):
                if not exp_files: 
                    st.error("Please upload files.")
                else:
                    st.session_state.pdf_files_res = None 
                    
                    with st.spinner("Extracting Export Data & Renaming PDFs..."):
                        extracted = extract_files(exp_files)
                        extracted.sort(key=lambda x: x.name)
                        
                        all_data = []
                        renamed_pdfs = [] 
                        
                        def format_date(d_str):
                            try: return datetime.strptime(d_str.strip(), "%Y/%m/%d")
                            except: return None
                        def to_float(v):
                            try: return round(float(re.sub(r'[^\d\.]', '', str(v))), 3)
                            except: return None
                        def to_int(v):
                            try: return int(re.sub(r'[^\d]', '', str(v)))
                            except: return None

                        for i, f_obj in enumerate(extracted, 1): 
                            f_obj.seek(0)
                            raw_pdf_bytes = f_obj.read() 
                            f_obj.seek(0)
                            
                            try:
                                reader = PdfReader(f_obj)
                                page1_text = reader.pages[0].extract_text() if len(reader.pages) > 0 else ""
                                full_text = " ".join([p.extract_text() for p in reader.pages])
                                
                                c_p1 = re.sub(r'\s+', ' ', page1_text)
                                c_f = re.sub(r'\s+', ' ', full_text)

                                decl_match = re.search(r"Declaration.*?(\d{12,15})", c_p1, re.IGNORECASE) or re.search(r"(\d{12,15})", c_p1)
                                decl_no = decl_match.group(1) if decl_match else ""
                                all_dates = re.findall(r"(\d{4}/\d{2}/\d{2})", c_f)
                                from collections import Counter
                                decl_date = format_date(Counter(all_dates).most_common(1)[0][0]) if all_dates else None
                                lic_match = re.search(r"(OVSEL\d+)", c_f)
                                license_no = lic_match.group(1) if lic_match else ""
                                inv_match = re.search(r"TOTAL\s*(\d+)\s*CTNS?\s*(\d+)\s*(?:U|PCS)\s*FOB\s*([\d\.]+)\s*CMP\s*([\d\.]+)", c_f, re.IGNORECASE)
                                t_ctn, t_pcs, fob_v, cmp_v = inv_match.groups() if inv_match else ("", "", "", "")
                                pack_match = re.search(r"TOTAL\s*\d+\s*CTNS?\s*\d+\s*(?:U|PCS)\s*([\d\.]+)\s*([\d\.]+)(?!\s*FOB)", c_f, re.IGNORECASE)
                                n_wt, g_wt = pack_match.groups() if pack_match else ("", "")

                                ctn_str = str(t_ctn).replace("/", "-").strip() if t_ctn else "0"
                                if "PK" not in ctn_str.upper():
                                    ctn_str += " PK"
                                
                                new_pdf_name = f"{i}_{decl_no}_{ctn_str}.pdf"
                                
                                final_pdf_bytes = raw_pdf_bytes
                                if do_compress and gs_setting:
                                    final_pdf_bytes = compress_pdf_bytes(final_pdf_bytes, gs_setting)
                                    
                                renamed_pdfs.append((new_pdf_name, final_pdf_bytes))

                                all_data.append({
                                    "No.": i, 
                                    "Declaration no.": decl_no, 
                                    "Declaration date": decl_date,
                                    "Export license no.": license_no, 
                                    "Total Pcs": to_int(t_pcs), 
                                    "Total Ctn": to_int(t_ctn),
                                    "Net weight": to_float(n_wt), 
                                    "Gross weight": to_float(g_wt), 
                                    "CMP (USD)": to_float(cmp_v), 
                                    "FOB (USD)": to_float(fob_v)
                                })
                            except: pass

                        if all_data:
                            df = pd.DataFrame(all_data)
                            out_xl = io.BytesIO()
                            with pd.ExcelWriter(out_xl, engine='openpyxl') as writer:
                                df.to_excel(writer, index=False, sheet_name='Export_Data', startrow=1)
                                ws = writer.sheets['Export_Data']
                                
                                ws.cell(row=1, column=1, value="TENG HUI (MYANMAR) KNITTING CO.,LTD_EXPORT DATA LISTS").font = Font(bold=True, size=16)
                                h_fill = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')
                                t_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
                                
                                for cell in ws[2]:
                                    cell.fill, cell.font, cell.border = h_fill, Font(bold=True), t_border
                                    cell.alignment = Alignment(horizontal='center', vertical='center')

                                last_row = len(df) + 2
                                for row in ws.iter_rows(min_row=3, max_row=last_row):
                                    for cell in row:
                                        cell.border = t_border
                                        cell.alignment = Alignment(vertical='center', horizontal='right')
                                        if cell.column in [1, 2, 3, 4]: cell.alignment = Alignment(horizontal='center')
                                        if cell.column == 3: cell.number_format = 'yyyy-mm-dd'

                                s_row = last_row + 1
                                ws.cell(row=s_row, column=4, value="TOTAL:").font = Font(bold=True)
                                for c_idx in range(5, 11):
                                    c_let = get_column_letter(c_idx)
                                    cell = ws.cell(row=s_row, column=c_idx, value=f"=SUM({c_let}3:{c_let}{last_row})")
                                    cell.font, cell.border = Font(bold=True), t_border
                                    cell.number_format = '#,##0' if c_idx in [5, 6] else '#,##0.00'

                                ws.auto_filter.ref = f"A2:{get_column_letter(df.shape[1])}{last_row}"
                                ws.freeze_panes = "A3"
                                
                                for col in ws.columns:
                                    c_let = get_column_letter(col[0].column)
                                    if c_let == 'A': ws.column_dimensions[c_let].width = 5.0
                                    else: ws.column_dimensions[c_let].width = max([len(str(c.value)) for c in col if c.value] + [0]) + 3

                            out_zip = io.BytesIO()
                            with zipfile.ZipFile(out_zip, "w", zipfile.ZIP_DEFLATED) as zf:
                                zf.writestr("All_Declarations_Summary.xlsx", out_xl.getvalue())
                                for pdf_name, pdf_bytes in renamed_pdfs:
                                    zf.writestr(pdf_name, pdf_bytes)
                            out_zip.seek(0)
                            
                            # 💡 Discord Alert
                            mm_time = datetime.utcnow() + timedelta(hours=6, minutes=30)
                            try:
                                send_discord_alert(st.session_state.current_user, "📄 **Export Data:** Successfully extracted Export Data & Renamed PDFs.", mm_time.strftime("%Y-%m-%d %H:%M:%S"))
                            except: pass

                            st.session_state.pdf_files_res = out_zip.getvalue()
                        else: st.warning("No data extracted.")

            pdf_data = st.session_state.get('pdf_files_res')
            if pdf_data is not None:
                if isinstance(pdf_data, bytes):
                    st.success("✅ Export Data Excel & Renamed PDFs successfully generated!")
                    st.download_button(
                        "📥 Download Excel & Renamed PDFs (ZIP)", 
                        data=pdf_data, 
                        file_name="Export_Data_and_Renamed_PDFs.zip", 
                        mime="application/zip", 
                        use_container_width=True
                    )
                else:
                    st.error("⚠️ System Warning: Data format is incorrect.")

            if c_pdf2.button("🗑️ Clear Files", key="c_pdf_btn", use_container_width=True): clear_files(); st.rerun()

        elif export_tool == "Rex Data":
            st.info("Upload multiple **PDFs** or **ZIP files** to extract REX Data into Excel.")
            rex_files = st.file_uploader("Upload REX PDFs or ZIP", type=["pdf", "zip"], accept_multiple_files=True, key=f"rex_{st.session_state.up_key}")
            
            c_rex1, c_rex2 = st.columns(2)
            if c_rex1.button("⚡ Extract Data", use_container_width=True, key="btn_rex_extract"):
                if not rex_files: st.error("Please upload files.")
                else:
                    with st.spinner("Extracting REX Data..."):
                        extracted = extract_files(rex_files)
                        all_data = []
                        
                        for f_obj in extracted:
                            f_obj.seek(0)
                            try:
                                with pdfplumber.open(f_obj) as pdf:
                                    text = "".join([p.extract_text() + "\n" for p in pdf.pages if p.extract_text()])
                                    ed_match = re.search(r'Declaration\s*No.*?(\d{12})', text, re.DOTALL | re.IGNORECASE) or re.search(r'\b(200\d{9})\b', text)
                                    el_ed_no = ed_match.group(1) if ed_match else ""

                                    inv_match = re.search(r'(TENG[\s\w\-]+EXP[\s\w\-/]+)', text, re.IGNORECASE) or re.search(r'Invoice\s*(?:No|No\.|Number)?\s*[:\s]*([A-Z0-9\-/]+)', text, re.IGNORECASE)
                                    inv_no = inv_match.group(1).strip().replace(" ", "").upper() if inv_match else ""

                                    dest_match = re.search(r'PORT OF DESTINATION\s*:\s*([A-Za-z\s]+)', text, re.IGNORECASE)
                                    dest = dest_match.group(1).strip() if dest_match else "UNITED KINGDOM"
                                    port_match = re.search(r'Warehouse\s+[A-Za-z0-9]+[\s\-]+([A-Za-z]+)', text, re.IGNORECASE)
                                    port = port_match.group(1).strip().upper() if port_match else "YANGON"

                                    raw_hs = list(re.finditer(r'\bHS\s*([\d\.]+)', text, re.IGNORECASE))
                                    valid_hs = [m for m in raw_hs if len(re.sub(r'\D', '', m.group(1))) >= 6]
                                    
                                    for i, match in enumerate(valid_hs):
                                        hs_code = re.sub(r'\D', '', match.group(1))[:6]
                                        end_pos = valid_hs[i+1].start() if i + 1 < len(valid_hs) else len(text)
                                        block = text[match.end():end_pos]
                                        
                                        n_match = re.search(r'Item name\s+(.*?)(?=\n|Quantity)', block, re.DOTALL | re.IGNORECASE)
                                        q_match = re.search(r'Quantity\s*\(\s*2\s*\)\D*([\d,\.]+)', block, re.IGNORECASE)
                                        v_match = re.search(r'Item value\D*([\d,\.]+)', block, re.IGNORECASE)
                                        
                                        if n_match and q_match and v_match:
                                            p_desc = n_match.group(1).strip()
                                            if p_desc.isdigit() or len(p_desc) < 3:
                                                 alt_name = re.search(r'([A-Za-z][A-Za-z\s%]+(?:Knitted|Woven)[^\n]+)', block)
                                                 if alt_name: p_desc = alt_name.group(1).strip()

                                            all_data.append({
                                                "Product Description": re.sub(r'\s+', ' ', p_desc).upper(),
                                                "HS (6-digit)": str(hs_code), "Destination": dest, "EL/ED No.": str(el_ed_no),
                                                "Invoice No": inv_no, "Port": port,
                                                "Quantity (PCS)": int(float(q_match.group(1).replace(',', ''))),
                                                "Value (USD)": float(v_match.group(1).replace(',', ''))
                                            })
                            except: pass

                        if all_data:
                            df = pd.DataFrame(all_data)
                            df.insert(0, 'No.', range(1, len(df) + 1))
                            df.insert(1, 'Company', "TENG HUI (MYANMAR) KNITTING CO.,LTD")
                            df.insert(2, 'Name REX Reg No.', "MMREX01032")
                            df.insert(5, 'Criteria', "W")
                            
                            cols = ['No.', 'Company', 'Name REX Reg No.', 'Product Description', 'HS (6-digit)', 'Criteria', 'Destination', 'EL/ED No.', 'Invoice No', 'Port', 'Quantity (PCS)', 'Value (USD)']
                            df = df.reindex(columns=cols)
                            df.columns = [c.upper() for c in df.columns]

                            out_xl = io.BytesIO()
                            with pd.ExcelWriter(out_xl, engine='openpyxl') as writer:
                                df.to_excel(writer, index=False, sheet_name='Sheet1')
                                ws = writer.sheets['Sheet1']
                                ws.auto_filter.ref = ws.dimensions
                                ws.freeze_panes = 'A2'
                                
                                for col in ws.columns:
                                    c_let = col[0].column_letter
                                    max_len = max([len(str(c.value)) for c in col if c.value] + [0])
                                    if c_let == 'D': ws.column_dimensions[c_let].width = 55
                                    elif c_let == 'B': ws.column_dimensions[c_let].width = 40
                                    elif c_let == 'I': ws.column_dimensions[c_let].width = 25
                                    else: ws.column_dimensions[c_let].width = max_len + 4
                                        
                                    for cell in col:
                                        if cell.row == 1:
                                            cell.font, cell.alignment = Font(bold=True), Alignment(horizontal='center', vertical='center')
                                        else:
                                            if c_let in ['K', 'L']: cell.alignment = Alignment(horizontal='right', vertical='center')
                                            elif c_let in ['B', 'D']: cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
                                            else: cell.alignment = Alignment(horizontal='center', vertical='center')
                                            
                                            if c_let in ['E', 'H']: cell.number_format = '@' 
                                            elif c_let == 'L': cell.number_format = '#,##0.00'

                            # 💡 Discord Alert
                            mm_time = datetime.utcnow() + timedelta(hours=6, minutes=30)
                            try:
                                send_discord_alert(st.session_state.current_user, "📋 **REX Data:** Successfully extracted REX Data to Excel.", mm_time.strftime("%Y-%m-%d %H:%M:%S"))
                            except: pass

                            st.session_state.rex_data_res = out_xl.getvalue()
                        else: st.warning("No data extracted.")

            if st.session_state.rex_data_res:
                st.success("✅ REX Data Excel successfully generated!")
                st.download_button("📥 Download REX Data (Excel)", st.session_state.rex_data_res, "Teng Hui_Rex_Data.xlsx", use_container_width=True)
            if c_rex2.button("🗑️ Clear Files", key="c_rex_btn", use_container_width=True): clear_files(); st.rerun()

        elif export_tool == "Extract Export Summary":
            st.info("Upload **Booking List (Excel)** to extract standard summary data. Items are grouped, PO Numbers are cleanly merged with commas, and errors are fixed.")
            
            if "export_summary_res" not in st.session_state: 
                st.session_state.export_summary_res = None
                
            export_summary_file = st.file_uploader("Upload Booking List (Excel)", type=["xlsx"], key=f"export_summary_{st.session_state.up_key}")
            
            c_exp_sum1, c_exp_sum2 = st.columns(2)
            if c_exp_sum1.button("⚡ Extract Standard Summary", use_container_width=True, key="btn_exp_sum_extract"):
                if not export_summary_file: 
                    st.error("❌ ကျေးဇူးပြု၍ Booking List Excel ဖိုင်ကို Upload တင်ပေးပါ။")
                else:
                    with st.spinner("Extracting and generating Standard Export Summary..."):
                        try:
                            xl = pd.ExcelFile(export_summary_file)
                            sheet_names = xl.sheet_names
                            
                            first_sheet_df = pd.read_excel(export_summary_file, sheet_name=sheet_names[0], nrows=0)
                            raw_title = first_sheet_df.columns[0] if len(first_sheet_df.columns) > 0 else ""
                            tod_match = re.search(r'TOD\s+(.*)', raw_title, re.IGNORECASE)
                            tod_val = tod_match.group(1).strip().upper() if tod_match else sheet_names[0]
                            
                            def clean_and_group_preserve_order(df):
                                df['WH CODE'] = df['WH CODE'].ffill()
                                df['COMPOSITION'] = df['COMPOSITION'].ffill()
                                df = df.dropna(subset=['DESCRIPTION'])
                                
                                if 'PO NO' in df.columns:
                                    df['PO NO'] = df['PO NO'].astype(str).str.replace(r'\.0$', '', regex=True).str.strip()
                                    df['PO NO'] = df['PO NO'].replace(['nan', 'NaN', 'None', ''], pd.NA).fillna("")
                                else:
                                    df['PO NO'] = ""
                                    
                                if 'HS CODE' in df.columns:
                                    df['HS CODE'] = df['HS CODE'].astype(str).str.replace(r'\.0$', '', regex=True).str.strip()
                                    df['HS CODE'] = df['HS CODE'].replace(['nan', 'NaN', 'None', ''], pd.NA).fillna("")
                                else:
                                    df['HS CODE'] = ""
                                
                                for col in ['PCS', 'CTNS', 'G.W', 'CBM']:
                                    df[col] = pd.to_numeric(df[col].astype(str).str.replace(',', ''), errors='coerce').fillna(0)
                                
                                df['orig_idx'] = df.index
                                
                                grouped = df.groupby(['WH CODE', 'DESCRIPTION', 'COMPOSITION', 'HS CODE'], as_index=False, dropna=False).agg({
                                    'PCS': 'sum',
                                    'CTNS': 'sum',
                                    'G.W': 'sum',
                                    'CBM': 'sum',
                                    'orig_idx': 'min'
                                })
                                
                                grouped = grouped.sort_values('orig_idx').drop(columns=['orig_idx'])
                                unique_wh_codes = df['WH CODE'].drop_duplicates().tolist()
                                
                                return grouped, unique_wh_codes

                            wb = openpyxl.Workbook()
                            wb.remove(wb.active)
                            
                            font_name = "Calibri"
                            header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                            header_font = Font(name=font_name, size=11, bold=True, color="FFFFFF")
                            subtotal_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
                            subtotal_font = Font(name=font_name, size=11, bold=True, color="000000")
                            grand_total_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
                            grand_total_font = Font(name=font_name, size=11, bold=True, color="000000")
                            regular_font = Font(name=font_name, size=11)
                            
                            center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
                            left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
                            right_align = Alignment(horizontal="right", vertical="center")
                            
                            thin_border = Side(style='thin', color='D9D9D9')
                            double_border = Side(style='double', color='000000')
                            top_thin_bottom_double = Border(top=thin_border, bottom=double_border, left=thin_border, right=thin_border)
                            cell_border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
                            
                            headers = ["WH CODE", "PO NO", "DESCRIPTION", "COMPOSITION", "Hs Code", "PCS", "CTNS", "G.W", "CBM"]
                            
                            processed_sheets = []
                            for s_name in sheet_names:
                                if 'TH' in s_name.upper() or 'CY' in s_name.upper():
                                    sheet_type = "TH Summary" if 'TH' in s_name.upper() else "CY Summary"
                                    
                                    df_orig = pd.read_excel(export_summary_file, sheet_name=s_name, header=1)
                                    df_orig['WH CODE'] = df_orig['WH CODE'].ffill()
                                    df_orig['COMPOSITION'] = df_orig['COMPOSITION'].ffill()
                                    
                                    df_clean_processing = df_orig.copy()
                                    if 'PO NO' in df_clean_processing.columns:
                                        df_clean_processing['PO NO'] = df_clean_processing['PO NO'].astype(str).str.replace(r'\.0$', '', regex=True).str.strip()
                                        df_clean_processing['PO NO'] = df_clean_processing['PO NO'].replace(['nan', 'NaN', 'None'], "")
                                    else:
                                        df_clean_processing['PO NO'] = ""
                                        
                                    if 'HS CODE' in df_clean_processing.columns:
                                        df_clean_processing['HS CODE'] = df_clean_processing['HS CODE'].astype(str).str.replace(r'\.0$', '', regex=True).str.strip()
                                        df_clean_processing['HS CODE'] = df_clean_processing['HS CODE'].replace(['nan', 'NaN', 'None'], "")
                                    else:
                                        df_clean_processing['HS CODE'] = ""
                                    
                                    def get_clean_pos(x):
                                        valid_pos = sorted(list(set([str(i) for i in x if str(i).strip() and str(i).lower() not in ['nan', 'none']])))
                                        return ", ".join(valid_pos)

                                    grouped_po_map = df_clean_processing.groupby(['WH CODE', 'DESCRIPTION', 'COMPOSITION', 'HS CODE'], dropna=False)['PO NO'].apply(get_clean_pos).to_dict()
                                    
                                    data, wh_order = clean_and_group_preserve_order(df_orig.copy())
                                    
                                    data['PO NO'] = data.set_index(['WH CODE', 'DESCRIPTION', 'COMPOSITION', 'HS CODE']).index.map(grouped_po_map).fillna("")
                                    
                                    processed_sheets.append((sheet_type, data, wh_order))
                                    
                            if not processed_sheets:
                                st.error("❌ TH သို့မဟုတ် CY ဆိုသည့် Sheet အမည်များကို ရှာမတွေ့ပါ။")
                            else:
                                for sheet_title, data, wh_order in processed_sheets:
                                    ws = wb.create_sheet(title=sheet_title)
                                    
                                    ws.append([f"H&M SHIPPING SUMMARY REPORT - TOD {tod_val}"])
                                    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
                                    ws.cell(row=1, column=1).font = Font(name=font_name, size=14, bold=True, color="1F4E78")
                                    ws.cell(row=1, column=1).alignment = Alignment(horizontal="left", vertical="center")
                                    ws.row_dimensions[1].height = 25
                                    ws.append([])
                                    
                                    ws.append(headers)
                                    header_row_idx = 3
                                    ws.row_dimensions[header_row_idx].height = 24
                                    for col_idx, cell in enumerate(ws[header_row_idx], 1):
                                        cell.fill = header_fill
                                        cell.font = header_font
                                        cell.alignment = center_align
                                        cell.border = cell_border

                                    ws.auto_filter.ref = f"A3:I3"
                                    ws.freeze_panes = "A4"

                                    current_row = 4
                                    
                                    for wh_code in wh_order:
                                        sub_df = data[data['WH CODE'] == wh_code]
                                        if sub_df.empty: continue
                                            
                                        start_group_row = current_row
                                        
                                        for _, row in sub_df.iterrows():
                                            ws.append([
                                                row['WH CODE'], row['PO NO'], row['DESCRIPTION'], row['COMPOSITION'], row['HS CODE'],
                                                row['PCS'], row['CTNS'], row['G.W'], row['CBM']
                                            ])
                                            ws.row_dimensions[current_row].height = 20
                                            for col_idx in range(1, 10):
                                                cell = ws.cell(row=current_row, column=col_idx)
                                                cell.font = regular_font
                                                cell.border = cell_border
                                                
                                                if col_idx in [1, 5]: 
                                                    cell.alignment = center_align
                                                elif col_idx in [2, 3, 4]: 
                                                    cell.alignment = left_align
                                                else: 
                                                    cell.alignment = right_align
                                                    
                                                if col_idx in [6, 7]: cell.number_format = '#,##0'
                                                elif col_idx in [8, 9]: cell.number_format = '#,##0.00'
                                            current_row += 1
                                            
                                        end_group_row = current_row - 1
                                        subtotal_label = f"Total {wh_code}"
                                        
                                        ws.append([
                                            subtotal_label, "", "", "", "",
                                            f"=SUM(F{start_group_row}:F{end_group_row})", 
                                            f"=SUM(G{start_group_row}:G{end_group_row})", 
                                            f"=SUM(H{start_group_row}:H{end_group_row})", 
                                            f"=SUM(I{start_group_row}:I{end_group_row})"
                                        ])
                                        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
                                        ws.row_dimensions[current_row].height = 22
                                        
                                        for col_idx in range(1, 10):
                                            cell = ws.cell(row=current_row, column=col_idx)
                                            cell.fill = subtotal_fill
                                            cell.font = subtotal_font
                                            cell.border = cell_border
                                            if col_idx == 1: cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
                                            elif col_idx >= 6:
                                                cell.alignment = right_align
                                                if col_idx in [6, 7]: cell.number_format = '#,##0'
                                                elif col_idx in [8, 9]: cell.number_format = '#,##0.00'
                                        
                                        current_row += 1
                                        ws.append([]) 
                                        current_row += 1
                                        
                                    ws.delete_rows(current_row - 1, 1)
                                    current_row -= 1
                                    
                                    ws.append([
                                        "GRAND TOTAL", "", "", "", "",
                                        f'=SUMIF($A$4:$A${current_row-1}, "Total*", F$4:F${current_row-1})', 
                                        f'=SUMIF($A$4:$A${current_row-1}, "Total*", G$4:G${current_row-1})', 
                                        f'=SUMIF($A$4:$A${current_row-1}, "Total*", H$4:H${current_row-1})', 
                                        f'=SUMIF($A$4:$A${current_row-1}, "Total*", I$4:I${current_row-1})'
                                    ])
                                    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
                                    ws.row_dimensions[current_row].height = 24
                                    
                                    for col_idx in range(1, 10):
                                        cell = ws.cell(row=current_row, column=col_idx)
                                        cell.fill = grand_total_fill
                                        cell.font = grand_total_font
                                        cell.border = top_thin_bottom_double
                                        if col_idx == 1: cell.alignment = Alignment(horizontal="left", vertical="center")
                                        elif col_idx >= 6:
                                            cell.alignment = right_align
                                            if col_idx in [6, 7]: cell.number_format = '#,##0'
                                            elif col_idx in [8, 9]: cell.number_format = '#,##0.00'
                                                
                                    col_widths = [18, 30, 35, 50, 15, 12, 12, 14, 14] 
                                    for idx, width in enumerate(col_widths, 1):
                                        ws.column_dimensions[get_column_letter(idx)].width = width
                                    ws.views.sheetView[0].showGridLines = True
                                    
                                out_xl = io.BytesIO()
                                wb.save(out_xl)
                                out_xl.seek(0)
                                
                                tod_format = tod_val.title()
                                output_filename = f"TOD {tod_format}_Export Summary.xlsx"
                                
                                # 💡 Discord Alert
                                mm_time = datetime.utcnow() + timedelta(hours=6, minutes=30)
                                try:
                                    send_discord_alert(st.session_state.current_user, f"📦 **Export Summary:** Generated perfect standard summary for TOD {tod_val}.", mm_time.strftime("%Y-%m-%d %H:%M:%S"))
                                except: pass

                                st.session_state.export_summary_res = {"name": output_filename, "data": out_xl.getvalue()}
                                
                        except Exception as e:
                            st.error(f"❌ Error processing files: {e}")
            
            if st.session_state.export_summary_res:
                res = st.session_state.export_summary_res
                st.success(f"✅ Successfully extracted Export Summary!")
                st.download_button("📥 Download Export Summary Excel", res["data"], res["name"], mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="btn_exp_sum_download", use_container_width=True)
                
            if c_exp_sum2.button("🗑️ Clear Files", key="c_exp_sum_clear", use_container_width=True):
                st.session_state.export_summary_res = None
                clear_files(); st.rerun()

        elif export_tool == "Extract Inv no":
            st.info("Upload **PDF** or **ZIP files** to extract H&M invoice data cleanly into a Standard Excel format.")
            
            if "hm_ext_res" not in st.session_state: st.session_state.hm_ext_res = None
            
            uploaded_hm_files = st.file_uploader("Upload PDF or ZIP", type=["pdf", "zip"], accept_multiple_files=True, key=f"hm_uploader_{st.session_state.up_key}")
            
            c_hm1, c_hm2 = st.columns(2)
            
            if c_hm1.button("⚡ Extract Data", key="btn_hm_extract", use_container_width=True):
                if not uploaded_hm_files: 
                    st.error("Please upload PDF or ZIP files.")
                else:
                    with st.spinner("Extracting data..."):
                        records = process_hm_files(uploaded_hm_files)
                        
                        if records:
                            df = pd.DataFrame(records)
                            output = io.BytesIO()
                            
                            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                                df.to_excel(writer, index=False, sheet_name='HM_Extracted_Data')
                                format_hm_excel(writer, df)
                                
                            # 💡 Discord Alert
                            mm_time = datetime.utcnow() + timedelta(hours=6, minutes=30)
                            try:
                                send_discord_alert(st.session_state.current_user, f"📝 **Extract INV no:** Extracted {len(records)} records.", mm_time.strftime("%Y-%m-%d %H:%M:%S"))
                            except: pass

                            st.session_state.hm_ext_res = {"count": len(records), "data": output.getvalue(), "df": df}
                        else:
                            st.warning("No records found in the uploaded files.")
                            
            if st.session_state.hm_ext_res:
                res = st.session_state.hm_ext_res
                st.success(f"Successfully extracted {res['count']} records!")
                st.dataframe(res["df"], use_container_width=True)
                
                st.download_button(
                    label="📥 Download Excel File",
                    data=res["data"],
                    file_name="HM_Invoice & Date_.xlsx", 
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="hm_download_btn",
                    use_container_width=True
                )
                
            if c_hm2.button("🗑️ Clear Files", key="c_hm_clear", use_container_width=True):
                st.session_state.hm_ext_res = None
                clear_files(); st.rerun()

    # ---------------------------------------------------------
    # 👥 USER RECORDS TAB
    # ---------------------------------------------------------
    with tab_records:
        st.write("### 📊 User Activity & Usage History")
        st.markdown("ဝက်ဘ်ဆိုက်အတွင်း အသုံးပြုသူများ၏ လုပ်ဆောင်ချက် မှတ်တမ်းများ")

        log_file = "activity_logs.csv"

        if os.path.exists(log_file):
            try:
                df = pd.read_csv(log_file)
                df = df.iloc[::-1]

                users_list = df['User'].dropna().unique().tolist() if 'User' in df.columns else []
                
                c_log1, c_log2, c_log3 = st.columns([2, 5, 1.5])
                with c_log1:
                    selected_user = st.selectbox("👤 Filter by User", ["All Users"] + users_list)
                with c_log3:
                    st.markdown("<div style='margin-top: 28px;'></div>", unsafe_allow_html=True)
                    if st.button("🔄 Refresh", use_container_width=True):
                        st.rerun()
                
                if selected_user != "All Users":
                    df = df[df['User'] == selected_user]

                st.dataframe(df, use_container_width=True, hide_index=True)
                
                st.download_button(
                    label="📥 Download Logs (CSV)",
                    data=df.to_csv(index=False).encode('utf-8'),
                    file_name=f"Activity_Logs_{datetime.now().strftime('%Y%m%d')}.csv",
                    mime="text/csv"
                )
                
            except Exception as e:
                st.error(f"မှတ်တမ်းများ ဖတ်ရှုရာတွင် အမှားအယွင်းရှိနေပါသည် - {e}")
        else:
            st.info("ယခုလောလောဆယ် မည်သည့် မှတ်တမ်းမျှ မရှိသေးပါ။")
