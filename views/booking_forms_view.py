import streamlit as st
import pandas as pd
import io
import os
import zipfile
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.drawing.image import Image as ExcelImage
from core.utils import add_log

# =========================================================
# HELPER: H&M BOOKING CONVERTER
# =========================================================
def convert_hm_booking(uploaded_file):
    output = io.BytesIO()
    import openpyxl
    wb = openpyxl.load_workbook(uploaded_file, read_only=True)
    visible_sheets = [sheet.title for sheet in wb.worksheets if sheet.sheet_state == 'visible']
    wb.close()
    
    xls = pd.ExcelFile(uploaded_file)
    all_rows = []
    
    main_title = ""
    if visible_sheets:
        try:
            df_title = pd.read_excel(uploaded_file, sheet_name=visible_sheets[0], header=None, nrows=1)
            main_title = str(df_title.iloc[0, 0])
            if main_title.lower() == 'nan': main_title = ""
        except: pass
            
    for sheet in visible_sheets:
        if sheet in xls.sheet_names:
            df = pd.read_excel(uploaded_file, sheet_name=sheet, header=1)
            if 'PO NO' in df.columns:
                all_rows.append(df)
            
    if not all_rows: return output
        
    combined_df = pd.concat(all_rows, ignore_index=True)
    combined_df['Group'] = combined_df['PO NO'].notna().cumsum()
    combined_df['PO NO'] = combined_df['PO NO'].ffill()
    combined_df['Depts'] = combined_df['Depts'].ffill()
    combined_df['SKU'] = combined_df['SKU'].ffill()
    combined_df['WH CODE'] = combined_df['WH CODE'].ffill()
    
    records = []; no_counter = 1 
    
    for grp, grp_df in combined_df.groupby('Group'):
        if grp_df.empty: continue
        first_row = grp_df.iloc[0]
        
        def safe_str(val): return str(val).strip() if pd.notna(val) and str(val).lower() != 'nan' else ""
        def safe_int_str(val):
            if pd.isna(val) or str(val).lower() == 'nan': return ""
            try: return str(int(val))
            except: return str(val).strip()
        
        po_no = safe_int_str(first_row['PO NO']).upper()
        wh_code = safe_str(first_row['WH CODE']).upper()
        ctns = first_row['CTNS'] if pd.notna(first_row['CTNS']) else 0
        gw = first_row['G.W'] if pd.notna(first_row['G.W']) else 0.0
        cbm = first_row['CBM'] if pd.notna(first_row['CBM']) else 0.0
        
        if not po_no or ctns == 0: continue
            
        desc_parts = []
        for _, row in grp_df.iterrows():
            d = safe_str(row['DESCRIPTION']).upper()
            if not d: continue
            c = safe_str(row['COMPOSITION']).upper()
            po = safe_int_str(row['PO NO']).upper()
            dept = safe_int_str(row['Depts']).upper()
            sku = safe_str(row['SKU']).upper()
            hs = safe_int_str(row['HS CODE']).upper()
            pcs = safe_int_str(row['PCS']).upper()
            
            sub_desc = [d]
            if c: sub_desc.append(c)
            sub_desc.append(f"PO NO: {po}-{dept}" if dept else f"PO NO: {po}")
            if sku: sub_desc.append(f"SKU: {sku}")
            if hs: sub_desc.append(f"HS CODE: {hs}")
            if pcs: sub_desc.append(f"QTY(PCS): {pcs}")
            desc_parts.append("\n".join(sub_desc))
            
        records.append({
            'NO.': no_counter, 'PO NO': po_no, 'DESCRIPTION': "\n\n".join(desc_parts),
            'WH CODE': wh_code, 'CTNS': int(ctns), 'G.W': gw, 'CBM': cbm
        })
        no_counter += 1
        
    res_df = pd.DataFrame(records)
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        if not res_df.empty:
            res_df.to_excel(writer, index=False, sheet_name='RESULT', startrow=1)
            ws = writer.sheets['RESULT']
            ws.sheet_view.showGridLines = False
            
            if main_title: ws['A1'] = main_title
            ws.merge_cells('A1:G1')
            ws['A1'].font = Font(size=14, bold=True)
            ws['A1'].alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[1].height = 30
            
            col_widths = {'A': 6, 'B': 15, 'C': 65, 'D': 15, 'E': 12, 'F': 12, 'G': 12}
            for col, width in col_widths.items(): ws.column_dimensions[col].width = width
                
            ws.freeze_panes = 'A3'
            ws.auto_filter.ref = f"A2:G{ws.max_row}"
            
            h_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            h_font = Font(color="FFFFFF", bold=True)
            z_fill1 = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
            z_fill2 = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
            a_cent = Alignment(horizontal="center", vertical="center", wrap_text=True)
            a_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
            a_right = Alignment(horizontal="right", vertical="center", wrap_text=True)
            t_bord = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
            
            for r_idx, row in enumerate(ws.iter_rows(min_row=2, max_col=7)):
                if r_idx == 0:
                    ws.row_dimensions[row[0].row].height = 25
                    for cell in row:
                        cell.fill, cell.font, cell.alignment, cell.border = h_fill, h_font, a_cent, t_bord
                else:
                    lines = (str(row[2].value).count('\n') + 1) if row[2].value else 1
                    ws.row_dimensions[row[0].row].height = max(40, lines * 16)
                    c_fill = z_fill1 if r_idx % 2 != 0 else z_fill2
                    for c_idx, cell in enumerate(row):
                        cell.fill, cell.border = c_fill, t_bord
                        if c_idx == 2: cell.alignment = a_left
                        elif c_idx in [4, 5, 6]: cell.alignment = a_right
                        else: cell.alignment = a_cent
                            
    output.seek(0)
    return output

# =========================================================
# MAIN UI
# =========================================================
def render_booking_forms_ui():
    st.subheader("📝 Booking Forms & Converters")
    
    if "bkg_up_key" not in st.session_state: st.session_state.bkg_up_key = 0
    if "air_bkg_res" not in st.session_state: st.session_state.air_bkg_res = None
        
    def clear_files(): 
        st.session_state.bkg_up_key += 1
        st.session_state.air_bkg_res = None

    selected_form = st.radio(
        "Select Booking Tool",
        ["Air Booking (Excel)", "Sea Booking (Format)"],
        horizontal=True, label_visibility="collapsed"
    )
    st.markdown("---")

    TEMPLATE_DIR = "assets/templates"

    # [1] Air Booking (Excel)
    if selected_form == "Air Booking (Excel)":
        st.markdown("#### ✈️ Air Booking Automation")
        # st.info အစား အောက်ပါကုဒ်ကို အသုံးပြုပါ
info_html = """
<div style="background-color: #EBF4FC; padding: 12px 16px; border-radius: 8px; display: flex; align-items: center;">
    <p style="margin: 0 !important; padding: 0 !important; color: #1E293B; line-height: 1.2;">
        Upload <b>H&M Air BKG Lists.xlsx</b> to generate automated Air Booking Excel forms.
    </p>
</div>
"""
st.markdown(info_html, unsafe_allow_html=True)
        bkg_file = st.file_uploader("Upload H&M Air BKG Lists (Excel)", type=["xlsx"], key=f"air_{st.session_state.bkg_up_key}")
        
        c_air1, c_air2 = st.columns(2)
        if c_air1.button("⚡ Generate Forms", width="stretch", type="primary"):
            if not bkg_file: 
                st.error("❌ ကျေးဇူးပြု၍ BKG List ဖိုင်ကို Upload တင်ပေးပါ။")
            else:
                with st.spinner("Generating Air Booking Excel files..."):
                    address_file = os.path.join(TEMPLATE_DIR, 'Addresses for HAWB [BKG] - 15.10.2025.xlsx')
                    logo_image_file = 'assets/images/logo.png'
                    stamp_image_file = 'assets/images/stamp.png' 
                    
                    if not os.path.exists(address_file):
                        st.error(f"❌ Error: '{address_file}' ကို Templates ဖိုဒါထဲတွင် ရှာမတွေ့ပါ။")
                    else:
                        try:
                            current_date = datetime.now().strftime("%d.%m.%y")
                            xl = pd.ExcelFile(bkg_file)
                            df_bkg = pd.read_excel(bkg_file, sheet_name=0, skiprows=1)
                            df_addr = pd.read_excel(address_file, sheet_name='HAWB')
                            tod_date = xl.sheet_names[0]
                            
                            light_grey = Side(style='thin', color='B0B0B0')
                            f_bord = Border(left=light_grey, right=light_grey, top=light_grey, bottom=light_grey)
                            d_bord = Border(right=light_grey, bottom=light_grey)

                            zip_buf = io.BytesIO(); file_count = 1
                            
                            with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
                                for _, row in df_bkg.iterrows():
                                    if pd.isna(row.get('WH CODE')): continue
                                    wh_code_bkg = str(row['WH CODE']).strip()
                                    order_no = str(row['PO NO']).split('.')[0] 
                                    depts_no = str(row['Depts']).split('.')[0] 
                                    country_code = str(row['SKU']).strip() if not pd.isna(row.get('SKU')) else ""
                                    
                                    mask = df_addr['Final Receiver'].str.contains(wh_code_bkg, na=False, case=False)
                                    addr_match = df_addr[mask]
                                    dest_name = addr_match.iloc[0]['Destination Country'] if not addr_match.empty else "NOT FOUND"
                                    consignee_val = addr_match.iloc[0]['HAWB Consignee '] if not addr_match.empty else "NOT FOUND"

                                    wb = Workbook()
                                    ws = wb.active
                                    ws.title = "Air Booking Form"
                                    ws.sheet_view.showGridLines = False

                                    ws.page_setup.paperSize = ws.PAPERSIZE_A4
                                    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
                                    ws.sheet_properties.pageSetUpPr.fitToPage = True
                                    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
                                    ws.page_margins.left = 0.25; ws.page_margins.right = 0.25
                                    ws.page_margins.top = 0.25; ws.page_margins.bottom = 0.25

                                    ws.column_dimensions['A'].width = 46; ws.column_dimensions['B'].width = 80
                                    ws.row_dimensions[1].height = 110
                                    ws.merge_cells('A1:B1')
                                    if os.path.exists(logo_image_file):
                                        ws.add_image(ExcelImage(logo_image_file), 'A1')
                                        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
                                        
                                    ws.row_dimensions[2].height = 25
                                    ws.merge_cells('A2:B2')
                                    ws['A2'] = f" DATE: {current_date}"
                                    ws['A2'].font, ws['A2'].alignment, ws['A2'].border = Font(size=10), Alignment(horizontal='left', vertical='center'), d_bord

                                    content = [
                                        ("SHIPPER NAME AND ADDRESS", "TENG HUI (MYANMAR) KNITTING COMPANY LIMITED\nNO.8/KHA KWIN NO.300/A, RAKHINE YOE GYI VILLAGE TRACT,\nHTAN TA BIN TOWNSHIP, YANGON REGION,\nREPUBLIC OF THE UNION OF MYANMAR."),
                                        ("CONSIGNEE NAME AND ADDRESS", consignee_val), ("NOTIFY", "SAME AS CONSIGNEE"), ("TOD DATE", tod_date),
                                        ("H & M ORDER NUMBER", f"{order_no}/{depts_no}"), ("DESCRIPTION OF CARGO", f"{row.get('DESCRIPTION', '')}\n{row.get('COMPOSITION', '')}"), 
                                        ("NUMBER OF PIECES (PCS)", row.get('PCS', '')), ("NUMBER OF CARTONS (CTN)", row.get('CTNS', '')),
                                        ("TOTAL GROSS WEIGHT (KG)", row.get('G.W', '')), ("TOTAL VOLUME (CBM)", row.get('CBM', '')),
                                        ("MODE OF TRANSPORT", "BY AIR"), ("DESTINATION AND COUNTRY CODE", f"{dest_name} / {country_code}"),
                                        ("FREIGHT TERM", "FREIGHT COLLECT"), ("SEASON AND DESTINATION WAREHOUSE", f"3-2026/{wh_code_bkg}"),
                                        ("H.S CODE", row.get('HS CODE', ''))
                                    ]

                                    current_row = 3
                                    for label, value in content:
                                        line_count = str(value).count('\n') + 1
                                        if "SHIPPER" in label: row_h = 85
                                        elif "CONSIGNEE" in label or "DESCRIPTION" in label: row_h = max(50, line_count * 18) 
                                        else: row_h = 30 
                                        ws.row_dimensions[current_row].height = row_h
                                        ca = ws.cell(row=current_row, column=1, value=f" {label}")
                                        ca.font, ca.border, ca.alignment = Font(size=10), f_bord, Alignment(horizontal='left', vertical='center')
                                        cb = ws.cell(row=current_row, column=2, value=value)
                                        cb.font, cb.border, cb.alignment = Font(size=10), f_bord, Alignment(wrap_text=True, horizontal='left', vertical='center')
                                        current_row += 1

                                    if os.path.exists(stamp_image_file):
                                        stamp = ExcelImage(stamp_image_file)
                                        stamp.width = 110; stamp.height = 110
                                        ws.add_image(stamp, f'B{current_row + 1}') 

                                    xlsx_buf = io.BytesIO()
                                    wb.save(xlsx_buf); xlsx_buf.seek(0)
                                    zf.writestr(f"{file_count:02d}. PO {order_no}_{country_code}_{wh_code_bkg}.xlsx", xlsx_buf.read())
                                    file_count += 1
                                    
                            if file_count > 1:
                                add_log(st.session_state.current_user, f"Generated {file_count-1} Air Booking forms")
                                st.session_state.air_bkg_res = {"count": file_count - 1, "zip_data": zip_buf.getvalue()}
                                
                        except Exception as e: st.error(f"❌ Error during processing: {e}")

        if st.session_state.air_bkg_res:
            res = st.session_state.air_bkg_res
            st.success(f"✅ Generated {res['count']} Excel files successfully!")
            st.download_button("📥 Download Generated Excel Files (ZIP)", res["zip_data"], "Air_Booking_Forms.zip", width="stretch")
            
        if c_air2.button("🗑️ Clear Files", width="stretch"): clear_files(); st.rerun()

    # [2] H&M Converter
    elif selected_form == "Sea Booking (Format)":
        st.markdown("#### 📝 H&M Booking List Converter")
        st.info("H&M SEA BOOKING LIST ဖိုင်ကို အောက်တွင် တင်ပေးပါ။")
        uploaded_hm_file = st.file_uploader("Upload Excel File", type=["xlsx"], key=f"hm_{st.session_state.bkg_up_key}")

        c_hm1, c_hm2 = st.columns(2)
        if c_hm1.button("⚡ Convert File", width="stretch", type="primary"):
            if not uploaded_hm_file: st.error("❌ ကျေးဇူးပြု၍ H&M Booking List ဖိုင်ကို အရင် Upload တင်ပေးပါ။")
            else:
                with st.spinner("Converting Data..."):
                    try:
                        converted_file = convert_hm_booking(uploaded_hm_file)
                        st.success("✅ အောင်မြင်စွာ ပြောင်းလဲပြီးပါပြီ!")
                        add_log(st.session_state.current_user, "Converted H&M Booking List")
                        st.download_button(
                            label="⬇️ Download Converted Result",
                            data=converted_file,
                            file_name=f"Converted_{uploaded_hm_file.name}",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            width="stretch"
                        )
                    except Exception as e: st.error(f"ဖိုင်ပြောင်းလဲရာတွင် အမှားအယွင်းဖြစ်ပေါ်နေပါသည်: {e}")
                        
        if c_hm2.button("🗑️ Clear Files", width="stretch"): clear_files(); st.rerun()
